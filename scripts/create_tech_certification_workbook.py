from __future__ import annotations

import argparse
import csv
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from common.data_paths import (  # noqa: E402
    KNOWN_COMPANY_DIR_TO_NAME,
    field_common_dir,
    normalize_field_name,
)
from data_intake.tech_intake.certification_standards import CERTIFICATION_SCHEMA  # noqa: E402

CERTIFICATION_LABELS = {
    "ISO_9001": "ISO 9001",
    "ISO_14001": "ISO 14001",
    "ISO_45001": "ISO 45001",
    "IATF_16949": "IATF 16949",
    "UL": "UL",
    "CE": "CE",
    "RoHS": "RoHS",
    "REACH": "REACH",
    "GMP": "GMP",
    "KGMP": "KGMP",
    "FDA": "FDA",
    "KC": "KC",
    "CUSTOMER_QUALITY_CERTIFICATION": "고객사 품질 인증",
    "SEMICONDUCTOR_QUALIFICATION": "반도체 고객사 qualification",
}

HEADERS = [
    "company_dir", "company", "ticker", "certification", "certification_key",
    "status", "year", "expiry_date", "source_type", "source_name", "source_url",
    "evidence_text", "directness", "confidence", "owner", "checked_at", "notes",
]

STATUS_VALUES = ["NEED_REVIEW", "CONFIRMED", "MENTIONED", "IN_PROGRESS", "EXPIRED_OR_REVOKED", "NOT_APPLICABLE"]
DIRECTNESS_VALUES = ["DIRECT_OR_COMPANY_DISCLOSED", "TEXT_MENTION", "CALCULATED", "UNKNOWN"]
CONFIDENCE_VALUES = ["HIGH", "MEDIUM", "LOW", "HIGH_NEGATIVE"]
SOURCE_TYPE_VALUES = ["manual", "company_homepage", "IR", "DART", "certification_body", "news", "customer_disclosure", "external_workbook"]


def _read_universe(path: Path | None) -> list[dict[str, str]]:
    if path is None or not path.exists():
        return [{"company_dir": k, "company": v, "ticker": ""} for k, v in KNOWN_COMPANY_DIR_TO_NAME.items()]
    rows: list[dict[str, str]] = []
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            slug = (row.get("company_dir") or row.get("slug") or row.get("company_slug") or row.get("dir") or "").strip()
            company = (row.get("company") or row.get("company_name") or row.get("name") or row.get("종목명") or row.get("기업명") or "").strip()
            ticker = (row.get("ticker") or row.get("stock_code") or row.get("종목코드") or "").strip()
            if not slug and company:
                # Keep a conservative filename-like fallback; data_paths can map many Korean names later.
                slug = company
            if slug or company:
                rows.append({"company_dir": slug, "company": company or slug, "ticker": ticker})
    return rows


def _style_workbook(wb: Any) -> None:
    try:
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.utils import get_column_letter
    except Exception as e:  # pragma: no cover
        raise RuntimeError("openpyxl이 필요합니다. pip install openpyxl 후 다시 실행하세요.") from e

    ws = wb["입력_인증근거"]
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    widths = {
        "A": 16, "B": 18, "C": 10, "D": 26, "E": 28, "F": 18,
        "G": 10, "H": 14, "I": 20, "J": 24, "K": 34, "L": 60,
        "M": 28, "N": 16, "O": 14, "P": 16, "Q": 34,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in range(2, ws.max_row + 1):
        ws.row_dimensions[row].height = 38
        for col in range(1, ws.max_column + 1):
            c = ws.cell(row, col)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = border

    dropdown = wb["드롭다운"]
    lists = {
        "F": ("status", len(STATUS_VALUES)),
        "I": ("source_type", len(SOURCE_TYPE_VALUES)),
        "M": ("directness", len(DIRECTNESS_VALUES)),
        "N": ("confidence", len(CONFIDENCE_VALUES)),
    }
    for col, (name, count) in lists.items():
        source_col = {"status": "A", "source_type": "B", "directness": "C", "confidence": "D"}[name]
        dv = DataValidation(type="list", formula1=f"=드롭다운!${source_col}$2:${source_col}${count+1}", allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{max(2, ws.max_row)}")

    # Highlight rows still requiring review.
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 6).fill = PatternFill("solid", fgColor="FFF2CC")

    for sheet_name in ["코드북_인증유형", "드롭다운", "README"]:
        sh = wb[sheet_name]
        for cell in sh[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col in range(1, sh.max_column + 1):
            sh.column_dimensions[get_column_letter(col)].width = 28
        for row in sh.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = border


def build_workbook(field: str, universe_csv: Path | None, output: Path) -> Path:
    try:
        from openpyxl import Workbook
    except Exception as e:  # pragma: no cover
        raise RuntimeError("openpyxl이 필요합니다. pip install openpyxl 후 다시 실행하세요.") from e

    companies = _read_universe(universe_csv)
    wb = Workbook()
    ws = wb.active
    ws.title = "입력_인증근거"
    ws.append(HEADERS)

    cert_items = [(k, CERTIFICATION_LABELS.get(k, k)) for k in CERTIFICATION_SCHEMA.keys()]
    for comp in companies:
        for cert_key, cert_label in cert_items:
            ws.append([
                comp.get("company_dir", ""), comp.get("company", ""), comp.get("ticker", ""),
                cert_label, cert_key, "NEED_REVIEW", "", "", "manual", "", "",
                "", "DIRECT_OR_COMPANY_DISCLOSED", "MEDIUM", "", "", "",
            ])

    codebook = wb.create_sheet("코드북_인증유형")
    codebook.append(["certification_key", "표시명", "category", "weight", "aliases", "해석"])
    for key, spec in CERTIFICATION_SCHEMA.items():
        codebook.append([
            key,
            CERTIFICATION_LABELS.get(key, key),
            spec.get("category", ""),
            spec.get("weight", ""),
            "; ".join(spec.get("aliases") or []),
            "직접 확인된 인증/표준은 상용화 준비도 보조 근거로 사용. 미확인은 감점하지 않음.",
        ])

    dd = wb.create_sheet("드롭다운")
    dd.append(["status", "source_type", "directness", "confidence"])
    max_len = max(len(STATUS_VALUES), len(SOURCE_TYPE_VALUES), len(DIRECTNESS_VALUES), len(CONFIDENCE_VALUES))
    for i in range(max_len):
        dd.append([
            STATUS_VALUES[i] if i < len(STATUS_VALUES) else "",
            SOURCE_TYPE_VALUES[i] if i < len(SOURCE_TYPE_VALUES) else "",
            DIRECTNESS_VALUES[i] if i < len(DIRECTNESS_VALUES) else "",
            CONFIDENCE_VALUES[i] if i < len(CONFIDENCE_VALUES) else "",
        ])

    readme = wb.create_sheet("README")
    readme.append(["항목", "설명"])
    readme.append(["사용법", "기업별 인증/표준을 확인한 뒤 status, source_name/url, evidence_text를 채우세요."])
    readme.append(["자동 로딩 경로", f"data\\{field}\\_sector_common\\tech_certifications\\certification_standards_external_signals.xlsx"])
    readme.append(["NEED_REVIEW", "자동 생성된 검토용 행입니다. evidence_text/source가 비어 있으면 data_intake가 근거로 사용하지 않습니다."])
    readme.append(["CONFIRMED", "공식 홈페이지, IR, DART, 인증기관 등에서 인증/승인 근거를 확인한 경우."])
    readme.append(["MENTIONED", "공식 인증서까지는 아니지만 문장 언급이 있는 경우."])
    readme.append(["IN_PROGRESS", "인증/qualification이 진행 중인 경우."])
    readme.append(["EXPIRED_OR_REVOKED", "만료·취소·실효 등 리스크 문구가 확인된 경우."])
    readme.append(["주의", "데이터 없음은 감점이 아니라 확인 제한입니다. 부정 근거가 있을 때만 리스크로 반영하세요."])
    readme.append(["created_at", datetime.now().isoformat(timespec="seconds")])

    _style_workbook(wb)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output


def main() -> int:
    parser = argparse.ArgumentParser(description="Create company-by-company certification standards Excel workbook for Tech Intake.")
    parser.add_argument("--field", default="반도체")
    parser.add_argument("--universe-csv", default="")
    parser.add_argument("--output", default="")
    args = parser.parse_args()

    field = normalize_field_name(args.field)
    universe = Path(args.universe_csv) if args.universe_csv else None
    output = Path(args.output) if args.output else field_common_dir("tech_certifications", field=field, create=True) / "certification_standards_external_signals.xlsx"
    out = build_workbook(field, universe, output)
    print(f"[OK] certification workbook created: {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
