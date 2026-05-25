from __future__ import annotations

r"""Check Valuation Agent outputs.

This script is intentionally runnable in two ways:

1) From the project root:
   python scripts\check_valuation_outputs.py --company-dir nepes

2) From any working directory:
   python C:\path\to\project\scripts\check_valuation_outputs.py --company-dir nepes

The project uses a src/ layout, so this file prepends <project>/src to sys.path
before importing common.data_paths. This prevents the common import error that can
occur when scripts are executed directly.
"""

import argparse
import json
import sys
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SRC_ROOT = PROJECT_ROOT / "src"

for _path in (PROJECT_ROOT, SRC_ROOT):
    _path_str = str(_path)
    if _path_str not in sys.path:
        sys.path.insert(0, _path_str)

try:
    from common.data_paths import company_agent_dir  # type: ignore
except Exception:
    company_agent_dir = None  # type: ignore

COMPANY_NAME_BY_SLUG = {
    "nepes": "네패스",
    "hanmi": "한미반도체",
    "hansol": "한솔케미칼",
    "duksan": "덕산테코피아",
    "ltc": "엘티씨",
}


def _read_json(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    for enc in ("utf-8-sig", "utf-8", "cp949"):
        try:
            return json.loads(path.read_text(encoding=enc))
        except UnicodeDecodeError:
            continue
        except Exception:
            return {}
    return {}


def _resolve_output_dir(company_dir: str, company_name: str | None = None) -> Path:
    if company_agent_dir is not None:
        try:
            return Path(company_agent_dir(company_dir, "valuation", create=False))
        except Exception:
            pass
    display_name = company_name or COMPANY_NAME_BY_SLUG.get(company_dir, company_dir)
    return PROJECT_ROOT / "data" / "반도체" / display_name / "valuation"


def _fmt(value: Any, digits: int = 2) -> str:
    if value is None or value == "":
        return "해당 없음"
    if isinstance(value, (int, float)):
        return f"{value:,.{digits}f}"
    return str(value)


def _scan_workbook_quality(path: Path) -> dict[str, Any]:
    result = {
        "exists": path.exists(),
        "sheet_count": 0,
        "new_polish_sheets": [],
        "problem_terms": {},
        "blank_like_cells": 0,
    }
    if not path.exists():
        return result
    try:
        import openpyxl  # type: ignore
        wb = openpyxl.load_workbook(path, data_only=False, read_only=True)
        result["sheet_count"] = len(wb.sheetnames)
        result["new_polish_sheets"] = [
            name for name in wb.sheetnames
            if name in {"26_투자판단_요약", "27_데이터보강_내역", "28_대시보드_연동키"}
        ]
        problem_terms = ["확인 제한", "원천 미제공", "추가 수집 필요", "분류 미제공", "진정한 가치평가"]
        counts = {term: 0 for term in problem_terms}
        blank_like = 0
        for ws in wb.worksheets:
            # Scan a bounded range to keep the check fast even for raw sheets.
            for row in ws.iter_rows(max_row=min(ws.max_row, 300), max_col=min(ws.max_column, 40)):
                row_has_value = any(cell.value not in (None, "") for cell in row)
                if not row_has_value:
                    continue
                for cell in row:
                    value = cell.value
                    if value in (None, ""):
                        blank_like += 1
                        continue
                    if isinstance(value, str):
                        for term in problem_terms:
                            if term in value:
                                counts[term] += 1
        result["problem_terms"] = counts
        result["blank_like_cells"] = blank_like
    except Exception as exc:
        result["error"] = str(exc)
    return result


def main() -> int:
    ap = argparse.ArgumentParser(description="Check Valuation Agent outputs")
    ap.add_argument("--company-dir", required=True, help="Company slug, e.g. nepes")
    ap.add_argument("--company", default=None, help="Optional display name, e.g. 네패스")
    args = ap.parse_args()

    out = _resolve_output_dir(args.company_dir, args.company)
    print(f"[CHECK] project_root={PROJECT_ROOT}")
    print(f"[CHECK] output_dir={out.resolve()}")

    files = [
        f"{args.company_dir}_valuation_metrics.json",
        f"{args.company_dir}_valuation_workbook.xlsx",
        f"{args.company_dir}_dashboard_payload.json",
        f"{args.company_dir}_valuation_validation.json",
        f"{args.company_dir}_valuation_report.md",
    ]
    missing_files: list[str] = []
    for name in files:
        p = out / name
        exists = p.exists()
        if not exists:
            missing_files.append(name)
        print(f"[CHECK] {name}: exists={exists} size={p.stat().st_size if exists else 0}")

    metrics = _read_json(out / f"{args.company_dir}_valuation_metrics.json")
    validation = _read_json(out / f"{args.company_dir}_valuation_validation.json")
    payload = _read_json(out / f"{args.company_dir}_dashboard_payload.json")

    dcf = metrics.get("dcf") or {}
    wacc = metrics.get("wacc") or {}
    price = metrics.get("price_summary") or {}
    peer = metrics.get("peer_comps") or {}
    ref = metrics.get("reference_universe_summary") or {}
    ml = metrics.get("ml_overlay") or {}
    advanced = metrics.get("advanced_valuation") or {}
    valuation_range = advanced.get("football_field") or {}

    status = validation.get("status") or metrics.get("validation_status") or "해당 없음"
    issues = validation.get("issues") or []

    print("[CHECK] validation_status=", status)
    print("[CHECK] opinion=", metrics.get("opinion"))
    print("[CHECK] wacc=", wacc.get("wacc"))
    print("[CHECK] implied_price=", dcf.get("implied_price"))
    print("[CHECK] enterprise_value=", dcf.get("enterprise_value"))
    print("[CHECK] equity_value=", dcf.get("equity_value"))
    print("[CHECK] price_rows=", price.get("price_rows"))
    print("[CHECK] latest_close=", price.get("latest_close"))
    print("[CHECK] market_cap=", price.get("market_cap"))
    print("[CHECK] shares_outstanding=", price.get("shares_outstanding"))
    print("[CHECK] avg_trading_value_20d=", price.get("avg_trading_value_20d"))
    print("[CHECK] target_psr=", peer.get("target_psr"))
    print("[CHECK] median_psr=", peer.get("median_psr"))
    print("[CHECK] psr_signal=", peer.get("psr_signal_kr") or peer.get("psr_signal"))
    print("[CHECK] reference_universe_rows=", ref.get("reference_universe_rows"))
    print("[CHECK] reference_focus_rows=", ref.get("reference_focus_rows"))
    print("[CHECK] target_peer_group=", ref.get("target_peer_group"))
    print("[CHECK] ml_signal=", ml.get("signal") or ml.get("signal_kr"))
    print("[CHECK] advanced_valuation_score=", advanced.get("advanced_valuation_score"))
    print("[CHECK] valuation_range_base_price=", valuation_range.get("base_price"))
    print("[CHECK] valuation_range_method_count=", valuation_range.get("method_count"))
    print("[CHECK] dashboard_payload_keys=", sorted(payload.keys()) if payload else "해당 없음")

    workbook_quality = _scan_workbook_quality(out / f"{args.company_dir}_valuation_workbook.xlsx")
    print("[CHECK] workbook_sheet_count=", workbook_quality.get("sheet_count"))
    print("[CHECK] workbook_polish_sheets=", ", ".join(workbook_quality.get("new_polish_sheets") or []) or "없음")
    print("[CHECK] workbook_problem_terms=", workbook_quality.get("problem_terms"))
    print("[CHECK] workbook_blank_like_cells_sample=", workbook_quality.get("blank_like_cells"))

    key_price_missing = []
    for key in ("latest_close", "market_cap", "shares_outstanding", "price_rows"):
        value = price.get(key)
        if value in (None, "", 0):
            key_price_missing.append(key)
    print("[CHECK] key_price_missing=", "없음" if not key_price_missing else ", ".join(key_price_missing))

    if issues:
        print("[CHECK] issues=")
        for issue in issues:
            print(f"  - {issue.get('severity')} / {issue.get('code')}: {issue.get('message')}")
    else:
        print("[CHECK] issues= 없음")

    if missing_files:
        print("[CHECK] missing_files=", ", ".join(missing_files))
        return 1
    if str(status).upper() not in {"PASS", "OK"}:
        return 1
    if key_price_missing:
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
