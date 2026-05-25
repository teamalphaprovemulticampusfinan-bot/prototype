from __future__ import annotations

"""Archive daily/monthly historical agent snapshots to Google Sheets.

History-only fast intake design
-------------------------------
This script is deliberately separated from the normal local pipeline.

Normal pipeline:
    agents may write JSON/MD under data/<field>/<company>/<agent>.

History pipeline:
    this script reads objective source files as of the requested day/month and
    uploads snapshots to Google Sheets.  It never writes under
    data/<field>/<company>/<agent>.

The goal is to avoid slow full-pipeline reruns while still making the historical
snapshot date-aware:
- market: reads data/market_excel and filters rows by company/date/month;
- issue: reads Issue_Integration / issue CSV/XLSX files and filters by company/date/month;
- macro: reads data/_global_common*/macro/macro_일별/월별/공통 files and filters by date/month;
- valuation: reads the existing valuation workbook and extracts as-of/month rows or formulas;
- finance/tech: uses existing local packets if available and annotates the requested cutoff.

All source filtering is deterministic.  No investment threshold or DMA weight is
defined in this file.
"""

import argparse
import csv
import json
import os
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterable

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SRC_DIR = PROJECT_ROOT / "src"
if str(SRC_DIR) not in os.sys.path:
    os.sys.path.insert(0, str(SRC_DIR))

try:
    from common.agent_history import save_agent_snapshot
except Exception as exc:  # pragma: no cover
    raise RuntimeError(
        "common.agent_history.save_agent_snapshot import failed. "
        "Run from project root and set PYTHONPATH to .\\src."
    ) from exc

AGENTS = ("finance", "market", "tech", "valuation", "issue", "macro")
DATE_COLUMNS = (
    "date",
    "날짜",
    "일자",
    "기준일",
    "기준일자",
    "as_of_date",
    "asof_date",
    "trading_date",
    "월말일",
    "month_end",
    "period_end",
)
COMPANY_COLUMNS = (
    "company",
    "회사명",
    "기업명",
    "종목명",
    "name",
    "company_name",
    "corp_name",
)
CODE_COLUMNS = (
    "stock_code",
    "ticker",
    "종목코드",
    "단축코드",
    "code",
    "symbol",
)


def _jsonable(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, dict):
        return {str(k): _jsonable(v) for k, v in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [_jsonable(v) for v in value]
    return str(value)


def _read_json(path: Path) -> dict[str, Any]:
    for enc in ("utf-8-sig", "utf-8", "cp949"):
        try:
            return json.loads(path.read_text(encoding=enc))
        except Exception:
            continue
    return {}


def _read_csv_rows(path: Path, *, max_rows: int | None = None) -> list[dict[str, Any]]:
    for enc in ("utf-8-sig", "utf-8", "cp949", "euc-kr"):
        try:
            with path.open("r", encoding=enc, newline="") as f:
                reader = csv.DictReader(f)
                rows = []
                for idx, row in enumerate(reader):
                    if max_rows is not None and idx >= max_rows:
                        break
                    rows.append({str(k).strip(): v for k, v in row.items()})
                return rows
        except Exception:
            continue
    return []


def _xlsx_rows(path: Path, *, max_rows: int | None = None) -> dict[str, list[dict[str, Any]]]:
    try:
        import openpyxl
    except Exception:
        return {}

    out: dict[str, list[dict[str, Any]]] = {}
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception:
        return out

    try:
        for ws in wb.worksheets:
            rows_iter = ws.iter_rows(values_only=True)
            try:
                header_raw = next(rows_iter)
            except StopIteration:
                continue
            headers = [str(h).strip() if h is not None else f"col_{i+1}" for i, h in enumerate(header_raw)]
            rows: list[dict[str, Any]] = []
            for idx, values in enumerate(rows_iter):
                if max_rows is not None and idx >= max_rows:
                    break
                row = {}
                for h, v in zip(headers, values):
                    row[h] = _jsonable(v)
                if any(v not in (None, "") for v in row.values()):
                    rows.append(row)
            if rows:
                out[ws.title] = rows
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return out


def _xlsx_key_values(path: Path, *, max_cells: int = 800) -> dict[str, Any]:
    """Extract a compact key/value view from an Excel workbook.

    This is used for valuation workbook snapshots.  It intentionally reads
    formulas' calculated values via data_only=True and does not modify the file.
    """
    try:
        import openpyxl
    except Exception:
        return {"status": "openpyxl_unavailable"}

    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        return {"status": "load_failed", "error": str(exc)}

    result: dict[str, Any] = {"status": "loaded", "sheets": {}}
    seen = 0
    try:
        for ws in wb.worksheets:
            sheet_items: list[dict[str, Any]] = []
            for row in ws.iter_rows(values_only=True):
                if seen >= max_cells:
                    break
                non_empty = [v for v in row if v not in (None, "")]
                if not non_empty:
                    continue
                # Common workbook pattern: label in col A, value in col B/C.
                key = str(row[0]).strip() if len(row) >= 1 and row[0] not in (None, "") else ""
                if key and len(row) >= 2:
                    sheet_items.append({"key": key, "value": _jsonable(row[1]), "extra": _jsonable(row[2:5])})
                    seen += 1
            if sheet_items:
                result["sheets"][ws.title] = sheet_items[:120]
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return result


def _parse_any_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, date) and not isinstance(value, datetime):
        return value
    if isinstance(value, datetime):
        return value.date()

    s = str(value).strip()
    if not s:
        return None
    s = s.replace(".", "-").replace("/", "-")
    s = re.sub(r"\s+.*$", "", s)
    # yyyymmdd
    if re.fullmatch(r"\d{8}", s):
        try:
            return datetime.strptime(s, "%Y%m%d").date()
        except Exception:
            return None
    # yyyy-mm
    if re.fullmatch(r"\d{4}-\d{1,2}", s):
        try:
            y, m = [int(x) for x in s.split("-")]
            # month rows are treated as month-end-ish date for comparison.
            import calendar
            return date(y, m, calendar.monthrange(y, m)[1])
        except Exception:
            return None
    # yyyy-mm-dd
    try:
        return datetime.strptime(s[:10], "%Y-%m-%d").date()
    except Exception:
        return None


def _date_from_row(row: dict[str, Any]) -> date | None:
    # exact column names first
    lowered = {str(k).strip().lower(): k for k in row}
    for c in DATE_COLUMNS:
        k = lowered.get(c.lower())
        if k is not None:
            d = _parse_any_date(row.get(k))
            if d:
                return d
    # loose contains
    for k, v in row.items():
        lk = str(k).strip().lower()
        if "date" in lk or "날짜" in lk or "일자" in lk or "기준" in lk:
            d = _parse_any_date(v)
            if d:
                return d
    return None


def _period_start(as_of_date: str, frequency: str) -> date:
    end = datetime.strptime(as_of_date, "%Y-%m-%d").date()
    if frequency == "monthly":
        return date(end.year, end.month, 1)
    return end


def _period_end(as_of_date: str) -> date:
    return datetime.strptime(as_of_date, "%Y-%m-%d").date()


def _row_period_match(row: dict[str, Any], *, as_of_date: str, frequency: str) -> tuple[bool, str]:
    """Return whether row belongs to the requested daily/monthly snapshot.

    If a source has no date column, the row is allowed but marked undated; the
    caller may still use company matching to avoid data loss.
    """
    d = _date_from_row(row)
    start = _period_start(as_of_date, frequency)
    end = _period_end(as_of_date)

    if d is None:
        return True, "undated_source_row"

    if frequency == "daily":
        if d == end:
            return True, "exact_daily_match"
        return False, "date_outside_daily"

    if start <= d <= end:
        return True, "within_month"
    return False, "date_outside_month"


def _norm_text(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "").strip()).lower()


def _digits(value: Any) -> str:
    return re.sub(r"\D+", "", str(value or ""))


def _row_text(row: dict[str, Any]) -> str:
    return " ".join(str(v) for v in row.values() if v not in (None, ""))


def _company_tokens(company: str, company_dir: str, stock_code: str = "", ticker: str = "") -> list[str]:
    tokens = [company, company_dir, stock_code, ticker]
    if stock_code and len(stock_code) < 6:
        tokens.append(stock_code.zfill(6))
    return [t for t in {_norm_text(x) for x in tokens if x} if t]


def _row_matches_company(row: dict[str, Any], *, company: str, company_dir: str, stock_code: str = "", ticker: str = "") -> bool:
    tokens = _company_tokens(company, company_dir, stock_code, ticker)
    if not tokens:
        return True

    # Prefer code columns when available.
    target_codes = {_digits(stock_code), _digits(ticker)}
    target_codes = {c.zfill(6) for c in target_codes if c}
    for k, v in row.items():
        lk = str(k).strip().lower()
        if lk in [c.lower() for c in CODE_COLUMNS] or "종목코드" in lk or "ticker" in lk or "code" in lk:
            rv = _digits(v)
            if rv and rv.zfill(6) in target_codes:
                return True

    text = _norm_text(_row_text(row))
    return any(t and t in text for t in tokens)


def _select_rows(
    rows: list[dict[str, Any]],
    *,
    company: str,
    company_dir: str,
    stock_code: str,
    ticker: str,
    as_of_date: str,
    frequency: str,
    limit: int = 80,
) -> dict[str, Any]:
    matched: list[dict[str, Any]] = []
    diagnostics = {
        "input_rows": len(rows),
        "matched_company_and_period": 0,
        "matched_company_only_fallback": 0,
        "selection_reason": "",
    }

    for row in rows:
        if not _row_matches_company(row, company=company, company_dir=company_dir, stock_code=stock_code, ticker=ticker):
            continue
        ok, reason = _row_period_match(row, as_of_date=as_of_date, frequency=frequency)
        if ok:
            r = dict(row)
            r["_history_match_reason"] = reason
            matched.append(r)

    if matched:
        diagnostics["matched_company_and_period"] = len(matched)
        diagnostics["selection_reason"] = "company_and_requested_period"
        return {"rows": matched[:limit], "diagnostics": diagnostics}

    # fallback: company match with latest <= as_of_date if date exists.
    end = _period_end(as_of_date)
    fallback: list[tuple[date | None, dict[str, Any]]] = []
    for row in rows:
        if not _row_matches_company(row, company=company, company_dir=company_dir, stock_code=stock_code, ticker=ticker):
            continue
        d = _date_from_row(row)
        if d is None or d <= end:
            fallback.append((d, row))

    fallback.sort(key=lambda x: x[0] or date.min, reverse=True)
    selected = []
    for d, row in fallback[:limit]:
        r = dict(row)
        r["_history_match_reason"] = "latest_company_row_before_as_of" if d else "undated_company_row_fallback"
        selected.append(r)

    diagnostics["matched_company_only_fallback"] = len(selected)
    diagnostics["selection_reason"] = "fallback_latest_company_rows_before_as_of" if selected else "no_matching_rows"
    return {"rows": selected, "diagnostics": diagnostics}


def _candidate_jsons(root: Path, *, field: str, company: str, company_dir: str, agent: str) -> list[Path]:
    bases = [
        root / "data" / field / company / agent,
        root / "data" / field / company_dir / agent,
        root / "data" / field / company / "auditor" / "first_auditor" / "compact_agent_packets",
    ]
    out: list[Path] = []
    for base in bases:
        if not base.exists():
            continue
        for p in base.rglob("*.json"):
            name = p.name.lower()
            if any(skip in name for skip in ("backup", "ablation_removed", ".bak")):
                continue
            if agent in name or "packet" in name or "summary" in name or "report" in name:
                out.append(p)
    return sorted(set(out), key=lambda p: p.stat().st_mtime, reverse=True)


def _load_base_agent_payload(
    *,
    field: str,
    company: str,
    company_dir: str,
    agent: str,
) -> tuple[dict[str, Any], str | None]:
    for path in _candidate_jsons(PROJECT_ROOT, field=field, company=company, company_dir=company_dir, agent=agent):
        data = _read_json(path)
        if data:
            return data, str(path.relative_to(PROJECT_ROOT))
    return {}, None


def _market_source_files(field: str, company: str, company_dir: str) -> list[Path]:
    roots = [
        PROJECT_ROOT / "data" / "market_excel",
        PROJECT_ROOT / "data" / field / "_sector_common",
        PROJECT_ROOT / "data" / field / "_sector_common" / "data",
        PROJECT_ROOT / "data" / field / company / "market",
        PROJECT_ROOT / "data" / field / company_dir / "market",
    ]
    out: list[Path] = []
    for root in roots:
        if not root.exists():
            continue
        for pat in ("*.xlsx", "*.xlsm", "*.csv"):
            out.extend(root.rglob(pat))
    return sorted(set(out), key=lambda p: p.stat().st_mtime, reverse=True)


def _issue_source_files(field: str, company: str, company_dir: str) -> list[Path]:
    roots = [
        PROJECT_ROOT / "data" / field / "_sector_common",
        PROJECT_ROOT / "data" / field / "_sector_common" / "data",
        PROJECT_ROOT / "data" / field / company / "issue",
        PROJECT_ROOT / "data" / field / company_dir / "issue",
        PROJECT_ROOT / "data" / "issue",
    ]
    out: list[Path] = []
    for root in roots:
        if not root.exists():
            continue
        for pat in ("*Issue*.xlsx", "*Issue*.csv", "*issue*.xlsx", "*issue*.csv", "*뉴스*.xlsx", "*뉴스*.csv", "*.rss.csv"):
            out.extend(root.rglob(pat))
    return sorted(set(out), key=lambda p: p.stat().st_mtime, reverse=True)


def _macro_source_files(frequency: str) -> list[Path]:
    roots = [
        PROJECT_ROOT / "data" / "_global_common" / "macro",
        PROJECT_ROOT / "data" / "_global_common_" / "macro",
    ]
    if frequency == "monthly":
        pats = ("macro_월별_*.csv", "macro_monthly_*.csv", "macro_공통_*.csv", "macro_common_*.csv")
    else:
        pats = ("macro_일별_*.csv", "macro_daily_*.csv", "macro_공통_*.csv", "macro_common_*.csv")
    out: list[Path] = []
    for root in roots:
        if not root.exists():
            continue
        for pat in pats:
            out.extend(root.glob(pat))
    return sorted(set(out), key=lambda p: p.stat().st_mtime, reverse=True)


def _valuation_workbooks(field: str, company: str, company_dir: str) -> list[Path]:
    roots = [
        PROJECT_ROOT / "data" / field / company / "valuation",
        PROJECT_ROOT / "data" / field / company_dir / "valuation",
    ]
    out: list[Path] = []
    for root in roots:
        if not root.exists():
            continue
        out.extend(root.glob("*_valuation_workbook.xlsx"))
        out.extend(root.glob("*valuation*.xlsx"))
    return sorted(set(out), key=lambda p: p.stat().st_mtime, reverse=True)


def _rows_from_file(path: Path, *, max_rows: int = 5000) -> dict[str, list[dict[str, Any]]]:
    if path.suffix.lower() == ".csv":
        return {path.stem: _read_csv_rows(path, max_rows=max_rows)}
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        return _xlsx_rows(path, max_rows=max_rows)
    return {}


def _enrich_tabular_sources(
    payload: dict[str, Any],
    *,
    source_name: str,
    files: list[Path],
    field: str,
    company: str,
    company_dir: str,
    stock_code: str,
    ticker: str,
    as_of_date: str,
    frequency: str,
    max_files: int = 8,
    max_rows_per_file: int = 80,
) -> dict[str, Any]:
    sources: list[dict[str, Any]] = []
    total_rows = 0

    for path in files[:max_files]:
        file_record: dict[str, Any] = {
            "path": str(path.relative_to(PROJECT_ROOT)) if path.is_relative_to(PROJECT_ROOT) else str(path),
            "mtime": datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds"),
            "sheets": {},
        }
        by_sheet = _rows_from_file(path)
        for sheet, rows in by_sheet.items():
            selected = _select_rows(
                rows,
                company=company,
                company_dir=company_dir,
                stock_code=stock_code,
                ticker=ticker,
                as_of_date=as_of_date,
                frequency=frequency,
                limit=max_rows_per_file,
            )
            if selected["rows"]:
                file_record["sheets"][sheet] = selected
                total_rows += len(selected["rows"])
        if file_record["sheets"]:
            sources.append(file_record)

    payload[f"history_{source_name}_source"] = {
        "mode": "history_fast_intake",
        "frequency": frequency,
        "as_of_date": as_of_date,
        "field": field,
        "company": company,
        "company_dir": company_dir,
        "stock_code": stock_code,
        "ticker": ticker,
        "source_file_count_scanned": len(files),
        "matched_source_count": len(sources),
        "matched_row_count": total_rows,
        "sources": sources,
    }
    return payload


def _enrich_market(
    payload: dict[str, Any],
    *,
    field: str,
    company: str,
    company_dir: str,
    stock_code: str,
    ticker: str,
    as_of_date: str,
    frequency: str,
) -> dict[str, Any]:
    return _enrich_tabular_sources(
        payload,
        source_name="market",
        files=_market_source_files(field, company, company_dir),
        field=field,
        company=company,
        company_dir=company_dir,
        stock_code=stock_code,
        ticker=ticker,
        as_of_date=as_of_date,
        frequency=frequency,
        max_files=10,
        max_rows_per_file=100,
    )




def _env_float(name: str, default: float) -> float:
    """Read a float environment variable with a safe fallback."""
    raw = os.environ.get(name)
    if raw is None or str(raw).strip() == "":
        return default
    try:
        value = float(str(raw).strip())
    except Exception:
        return default
    if not (value >= 0):
        return default
    return value


def _env_int(name: str, default: int) -> int:
    """Read an integer environment variable with a safe fallback."""
    raw = os.environ.get(name)
    if raw is None or str(raw).strip() == "":
        return default
    try:
        value = int(float(str(raw).strip()))
    except Exception:
        return default
    return value if value >= 0 else default


def _previous_month_window(as_of_date: str) -> tuple[date, date]:
    """Return the full previous calendar-month window for a monthly issue lag."""
    end = _period_end(as_of_date)
    cur_start = date(end.year, end.month, 1)
    lag_end = cur_start - timedelta(days=1)
    lag_start = date(lag_end.year, lag_end.month, 1)
    return lag_start, lag_end


def _issue_time_windows(as_of_date: str, frequency: str) -> dict[str, tuple[date, date]]:
    """Return current and lag windows for history-only issue intake.

    Monthly:
      current = current calendar month until as_of_date
      lag     = immediately previous calendar month

    Daily:
      current = as_of_date only
      lag     = previous N calendar days, excluding as_of_date
    """
    end = _period_end(as_of_date)
    if frequency == "monthly":
        return {
            "current": (_period_start(as_of_date, frequency), end),
            "lag": _previous_month_window(as_of_date),
        }

    lag_days = _env_int("ALPHAPROVE_ISSUE_DAILY_LAG_DAYS", 30)
    lag_end = end - timedelta(days=1)
    lag_start = lag_end - timedelta(days=max(lag_days - 1, 0))
    return {
        "current": (end, end),
        "lag": (lag_start, lag_end),
    }


def _issue_window_weights() -> tuple[float, float]:
    """Return source-window weights for issue evidence.

    These are not DMA posterior weights and are not investment thresholds.
    They only annotate how much current-period and lag-period issue evidence
    should contribute when the history snapshot is summarized.
    """
    current = _env_float("ALPHAPROVE_ISSUE_CURRENT_WINDOW_WEIGHT", 0.70)
    lag = _env_float("ALPHAPROVE_ISSUE_LAG_WINDOW_WEIGHT", 0.30)
    total = current + lag
    if total <= 0:
        return 0.70, 0.30
    return current / total, lag / total


def _select_rows_for_explicit_window(
    rows: list[dict[str, Any]],
    *,
    company: str,
    company_dir: str,
    stock_code: str,
    ticker: str,
    window_start: date,
    window_end: date,
    bucket: str,
    weight: float,
    limit: int,
) -> dict[str, Any]:
    """Select company rows inside an explicit issue time window."""
    matched: list[tuple[date, dict[str, Any]]] = []
    company_match_count = 0
    dated_company_rows = 0

    for row in rows:
        if not _row_matches_company(row, company=company, company_dir=company_dir, stock_code=stock_code, ticker=ticker):
            continue
        company_match_count += 1
        d = _date_from_row(row)
        if d is None:
            continue
        dated_company_rows += 1
        if window_start <= d <= window_end:
            matched.append((d, row))

    matched.sort(key=lambda x: x[0], reverse=True)

    selected: list[dict[str, Any]] = []
    for d, row in matched[:limit]:
        r = dict(row)
        r["_history_match_reason"] = f"issue_{bucket}_window_match"
        r["_issue_time_bucket"] = bucket
        r["_issue_time_weight"] = weight
        r["_issue_window_start"] = window_start.isoformat()
        r["_issue_window_end"] = window_end.isoformat()
        r["_issue_weight_role"] = "source_window_weight_not_dma_weight"
        r["_issue_row_date"] = d.isoformat()
        selected.append(r)

    diagnostics = {
        "input_rows": len(rows),
        "company_match_count": company_match_count,
        "dated_company_rows": dated_company_rows,
        "matched_rows": len(selected),
        "window_start": window_start.isoformat(),
        "window_end": window_end.isoformat(),
        "bucket": bucket,
        "source_window_weight": weight,
        "selection_reason": f"issue_{bucket}_explicit_window",
    }
    return {"rows": selected, "diagnostics": diagnostics}

def _enrich_issue(
    payload: dict[str, Any],
    *,
    field: str,
    company: str,
    company_dir: str,
    stock_code: str,
    ticker: str,
    as_of_date: str,
    frequency: str,
) -> dict[str, Any]:
    """History-only issue intake with current-window + lag-window evidence.

    This function is used only while creating Google Sheets history snapshots.
    It does not run the normal local pipeline and does not write to
    data/<field>/<company>/issue.

    Monthly default:
      current month issue rows: 70%
      previous month issue rows: 30%

    Daily default:
      current day issue rows: 70%
      previous 30 calendar days, excluding current day: 30%

    The 70/30 values are source-window evidence weights.  They are deliberately
    separated from DMA posterior weights and from final recommendation logic.
    """
    current_weight, lag_weight = _issue_window_weights()
    windows = _issue_time_windows(as_of_date, frequency)

    files = _issue_source_files(field, company, company_dir)
    max_files = 12
    max_rows_per_file = 120
    current_limit = max(1, int(round(max_rows_per_file * current_weight)))
    lag_limit = max(1, max_rows_per_file - current_limit)

    sources: list[dict[str, Any]] = []
    total_current_rows = 0
    total_lag_rows = 0
    total_weighted_rows = 0.0

    for path in files[:max_files]:
        file_record: dict[str, Any] = {
            "path": str(path.relative_to(PROJECT_ROOT)) if path.is_relative_to(PROJECT_ROOT) else str(path),
            "mtime": datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds"),
            "sheets": {},
        }
        by_sheet = _rows_from_file(path)

        for sheet, rows in by_sheet.items():
            cur_start, cur_end = windows["current"]
            lag_start, lag_end = windows["lag"]

            current_selected = _select_rows_for_explicit_window(
                rows,
                company=company,
                company_dir=company_dir,
                stock_code=stock_code,
                ticker=ticker,
                window_start=cur_start,
                window_end=cur_end,
                bucket="current",
                weight=current_weight,
                limit=current_limit,
            )
            lag_selected = _select_rows_for_explicit_window(
                rows,
                company=company,
                company_dir=company_dir,
                stock_code=stock_code,
                ticker=ticker,
                window_start=lag_start,
                window_end=lag_end,
                bucket="lag",
                weight=lag_weight,
                limit=lag_limit,
            )

            combined_rows = current_selected["rows"] + lag_selected["rows"]
            if not combined_rows:
                continue

            weighted_count = (
                len(current_selected["rows"]) * current_weight
                + len(lag_selected["rows"]) * lag_weight
            )
            sheet_record = {
                "rows": combined_rows,
                "current": current_selected,
                "lag": lag_selected,
                "diagnostics": {
                    "selection_reason": "issue_current_window_plus_lag_window",
                    "current_rows": len(current_selected["rows"]),
                    "lag_rows": len(lag_selected["rows"]),
                    "weighted_row_count": weighted_count,
                    "current_window": {
                        "start": cur_start.isoformat(),
                        "end": cur_end.isoformat(),
                        "weight": current_weight,
                    },
                    "lag_window": {
                        "start": lag_start.isoformat(),
                        "end": lag_end.isoformat(),
                        "weight": lag_weight,
                    },
                },
            }

            file_record["sheets"][sheet] = sheet_record
            total_current_rows += len(current_selected["rows"])
            total_lag_rows += len(lag_selected["rows"])
            total_weighted_rows += weighted_count

        if file_record["sheets"]:
            sources.append(file_record)

    payload["history_issue_lag_policy"] = {
        "mode": "history_fast_intake_issue_current_plus_lag",
        "frequency": frequency,
        "as_of_date": as_of_date,
        "current_window_weight": current_weight,
        "lag_window_weight": lag_weight,
        "daily_lag_days": _env_int("ALPHAPROVE_ISSUE_DAILY_LAG_DAYS", 30),
        "weight_role": "source_evidence_window_weight_not_dma_posterior_weight",
        "dma_unchanged": True,
        "final_recommendation_threshold_unchanged": True,
    }
    payload["history_issue_source"] = {
        "mode": "history_fast_intake",
        "frequency": frequency,
        "as_of_date": as_of_date,
        "field": field,
        "company": company,
        "company_dir": company_dir,
        "stock_code": stock_code,
        "ticker": ticker,
        "source_file_count_scanned": len(files),
        "matched_source_count": len(sources),
        "current_row_count": total_current_rows,
        "lag_row_count": total_lag_rows,
        "matched_row_count": total_current_rows + total_lag_rows,
        "weighted_row_count": total_weighted_rows,
        "sources": sources,
    }
    return payload



def _enrich_macro(
    payload: dict[str, Any],
    *,
    as_of_date: str,
    frequency: str,
) -> dict[str, Any]:
    files = _macro_source_files(frequency)
    sources: list[dict[str, Any]] = []
    total = 0
    for path in files[:8]:
        rows = _read_csv_rows(path)
        selected: list[dict[str, Any]] = []
        for row in rows:
            ok, reason = _row_period_match(row, as_of_date=as_of_date, frequency=frequency)
            if ok:
                r = dict(row)
                r["_history_match_reason"] = reason
                selected.append(r)

        if not selected:
            # objective fallback: latest row <= as_of_date
            end = _period_end(as_of_date)
            pairs = []
            for row in rows:
                d = _date_from_row(row)
                if d and d <= end:
                    pairs.append((d, row))
            pairs.sort(key=lambda x: x[0], reverse=True)
            if pairs:
                r = dict(pairs[0][1])
                r["_history_match_reason"] = "latest_macro_row_before_as_of"
                selected = [r]

        if selected:
            sources.append(
                {
                    "path": str(path.relative_to(PROJECT_ROOT)) if path.is_relative_to(PROJECT_ROOT) else str(path),
                    "mtime": datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds"),
                    "rows": selected[:40],
                    "matched_row_count": len(selected),
                }
            )
            total += len(selected)

    payload["history_macro_source"] = {
        "mode": "history_fast_intake",
        "frequency": frequency,
        "as_of_date": as_of_date,
        "source_file_count_scanned": len(files),
        "matched_source_count": len(sources),
        "matched_row_count": total,
        "sources": sources,
    }
    return payload


def _enrich_valuation(
    payload: dict[str, Any],
    *,
    field: str,
    company: str,
    company_dir: str,
    stock_code: str,
    ticker: str,
    as_of_date: str,
    frequency: str,
) -> dict[str, Any]:
    files = _valuation_workbooks(field, company, company_dir)
    sources = []
    for path in files[:3]:
        source = {
            "path": str(path.relative_to(PROJECT_ROOT)) if path.is_relative_to(PROJECT_ROOT) else str(path),
            "mtime": datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds"),
            "formula_values": _xlsx_key_values(path),
        }
        by_sheet = _xlsx_rows(path, max_rows=3000)
        matched_sheets = {}
        for sheet, rows in by_sheet.items():
            selected = _select_rows(
                rows,
                company=company,
                company_dir=company_dir,
                stock_code=stock_code,
                ticker=ticker,
                as_of_date=as_of_date,
                frequency=frequency,
                limit=80,
            )
            if selected["rows"]:
                matched_sheets[sheet] = selected
        if matched_sheets:
            source["matched_rows"] = matched_sheets
        sources.append(source)

    payload["history_valuation_source"] = {
        "mode": "history_fast_intake",
        "frequency": frequency,
        "as_of_date": as_of_date,
        "source_file_count_scanned": len(files),
        "matched_source_count": len(sources),
        "sources": sources,
    }
    return payload


def build_history_snapshot(
    *,
    field: str,
    company: str,
    company_dir: str,
    agent: str,
    as_of_date: str,
    frequency: str,
    stock_code: str = "",
    ticker: str = "",
    history_fast_intake: bool = True,
) -> tuple[dict[str, Any], str | None]:
    """Build one history snapshot payload for one company/agent."""
    payload, source_file = _load_base_agent_payload(field=field, company=company, company_dir=company_dir, agent=agent)

    if not isinstance(payload, dict):
        payload = {}

    payload.setdefault("agent", agent)
    payload.setdefault("company", company)
    payload.setdefault("company_dir", company_dir)
    if stock_code:
        payload.setdefault("stock_code", stock_code)
    if ticker:
        payload.setdefault("ticker", ticker)

    payload["history_snapshot_meta"] = {
        "mode": "history_fast_intake" if history_fast_intake else "history_base_snapshot",
        "field": field,
        "company": company,
        "company_dir": company_dir,
        "agent": agent,
        "as_of_date": as_of_date,
        "frequency": frequency,
        "period_start": _period_start(as_of_date, frequency).isoformat(),
        "period_end": _period_end(as_of_date).isoformat(),
        "base_payload_source_file": source_file,
        "local_agent_folder_write": False,
        "normal_pipeline_unchanged": True,
    }

    if history_fast_intake:
        if agent == "market":
            payload = _enrich_market(
                payload,
                field=field,
                company=company,
                company_dir=company_dir,
                stock_code=stock_code,
                ticker=ticker,
                as_of_date=as_of_date,
                frequency=frequency,
            )
        elif agent == "issue":
            payload = _enrich_issue(
                payload,
                field=field,
                company=company,
                company_dir=company_dir,
                stock_code=stock_code,
                ticker=ticker,
                as_of_date=as_of_date,
                frequency=frequency,
            )
        elif agent == "macro":
            payload = _enrich_macro(payload, as_of_date=as_of_date, frequency=frequency)
        elif agent == "valuation":
            payload = _enrich_valuation(
                payload,
                field=field,
                company=company,
                company_dir=company_dir,
                stock_code=stock_code,
                ticker=ticker,
                as_of_date=as_of_date,
                frequency=frequency,
            )

    return _jsonable(payload), source_file


def _read_universe(path: Path | None, *, default_five: bool = False) -> list[dict[str, str]]:
    if path and path.exists():
        rows = _read_csv_rows(path)
        out = []
        for row in rows:
            company_dir = (
                row.get("company_dir")
                or row.get("slug")
                or row.get("company_slug")
                or row.get("기업폴더")
                or row.get("폴더")
                or ""
            )
            company = (
                row.get("company")
                or row.get("company_name")
                or row.get("name")
                or row.get("기업명")
                or row.get("회사명")
                or ""
            )
            stock_code = (
                row.get("stock_code")
                or row.get("ticker")
                or row.get("종목코드")
                or row.get("단축코드")
                or row.get("code")
                or ""
            )
            ticker = row.get("ticker") or stock_code
            if company_dir and company:
                out.append(
                    {
                        "company_dir": str(company_dir).strip(),
                        "company": str(company).strip(),
                        "stock_code": str(stock_code).strip().zfill(6) if str(stock_code).strip().isdigit() else str(stock_code).strip(),
                        "ticker": str(ticker).strip().zfill(6) if str(ticker).strip().isdigit() else str(ticker).strip(),
                    }
                )
        if out:
            return out

    if default_five:
        return [
            {"company_dir": "nepes", "company": "네패스", "stock_code": "033640", "ticker": "033640"},
            {"company_dir": "hanmi", "company": "한미반도체", "stock_code": "042700", "ticker": "042700"},
            {"company_dir": "hansol", "company": "한솔케미칼", "stock_code": "014680", "ticker": "014680"},
            {"company_dir": "duksan", "company": "덕산테코피아", "stock_code": "317330", "ticker": "317330"},
            {"company_dir": "ltc", "company": "엘티씨", "stock_code": "170920", "ticker": "170920"},
        ]

    raise FileNotFoundError(f"universe csv not found or empty: {path}")


def archive_snapshots(
    *,
    field: str,
    as_of_date: str,
    frequency: str,
    companies: list[dict[str, str]],
    agents: Iterable[str] = AGENTS,
    overwrite: bool = True,
    history_fast_intake: bool = True,
) -> list[dict[str, Any]]:
    """Archive snapshots to Google Sheets only.

    No files are written to data/<field>/<company>/<agent>.
    """
    results: list[dict[str, Any]] = []
    for item in companies:
        company_dir = item["company_dir"]
        company = item["company"]
        stock_code = item.get("stock_code", "")
        ticker = item.get("ticker", stock_code)
        print(f"[archive] {as_of_date} {company}/{company_dir} snapshots -> Google Sheets only")

        for agent in agents:
            agent = agent.strip().lower()
            if not agent:
                continue
            payload, source_file = build_history_snapshot(
                field=field,
                company=company,
                company_dir=company_dir,
                agent=agent,
                as_of_date=as_of_date,
                frequency=frequency,
                stock_code=stock_code,
                ticker=ticker,
                history_fast_intake=history_fast_intake,
            )

            res = save_agent_snapshot(
                as_of_date=as_of_date,
                field=field,
                company_dir=company_dir,
                company_name=company,
                agent=agent,
                payload=payload,
                source_file=source_file,
                source_mode=f"history_{frequency}_fast_intake_google_sheets_only",
                overwrite=overwrite,
            )
            results.append({"company_dir": company_dir, "company": company, "agent": agent, "result": res})
            print(f"  - {agent}: saved source={source_file or 'generated_history_snapshot'}")

    return results


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Archive date-aware history snapshots to Google Sheets only.")
    parser.add_argument("--field", default="반도체")
    parser.add_argument("--as-of-date", required=True)
    parser.add_argument("--frequency", choices=["daily", "monthly"], required=True)
    parser.add_argument("--universe-csv")
    parser.add_argument("--default-five", action="store_true")
    parser.add_argument("--agents", default=",".join(AGENTS))
    parser.add_argument("--no-overwrite", action="store_true")
    parser.add_argument("--no-history-fast-intake", action="store_true")
    args = parser.parse_args(argv)

    companies = _read_universe(Path(args.universe_csv) if args.universe_csv else None, default_five=args.default_five)
    agents = [a.strip().lower() for a in args.agents.replace(";", ",").split(",") if a.strip()]
    archive_snapshots(
        field=args.field,
        as_of_date=args.as_of_date,
        frequency=args.frequency,
        companies=companies,
        agents=agents,
        overwrite=not args.no_overwrite,
        history_fast_intake=not args.no_history_fast_intake,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
