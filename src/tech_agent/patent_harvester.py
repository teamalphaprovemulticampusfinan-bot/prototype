from __future__ import annotations

import csv
import json
import re
from collections import Counter
from datetime import datetime
from pathlib import Path
from common.data_paths import company_agent_dir, company_common_dir, company_config_path, field_agent_dir, field_common_dir, ml_universe_dir, tech_source_dir
from typing import Any

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None

try:
    from .config import ROOT_DIR, OUTPUT_DIR
except Exception:
    ROOT_DIR = Path(__file__).resolve().parents[2]
    OUTPUT_DIR = field_agent_dir("tech")

try:
    from .utils import clean_text, ensure_dir
except Exception:
    def clean_text(value: Any) -> str:
        if value is None:
            return ""
        return re.sub(r"\s+", " ", str(value)).strip()

    def ensure_dir(path: str | Path) -> Path:
        p = Path(path)
        p.mkdir(parents=True, exist_ok=True)
        return p


PATENT_EXTS = {".csv", ".tsv", ".xlsx", ".xlsm", ".txt", ".md", ".json"}
PATENT_FILE_KEYWORDS = ["kipris", "patent", "patents", "특허", "산업재산권", "bibliographic"]
GENERATED_PATENT_SKIP_KEYWORDS = [
    "claims_normalized",
    "citations_normalized",
    "family_normalized",
    "tech_ip_evidence_composite",
    "tech_ip_claim_features",
    "tech_ip_citation_features",
    "tech_ip_family_features",
    "tech_ip_legal_features",
    "tech_ip_ml",
    "tech_peer",
    "excel_quantified",
    "template",
    "inventory",
    "summary",
    "request_targets",
    "similarity",
    "top_pairs",
    "peer",
]

SEMICONDUCTOR_IPC_PREFIXES = {
    "H01L", "H10B", "H10D", "H10F", "H10H", "H10K", "H10N", "H10P", "H10W",
    "G01R", "G06F", "G06N", "B81B", "B81C", "C23C", "C30B",
}

TECH_BASE_TERMS = [
    "WLP", "FOWLP", "FO-WLP", "FOPLP", "FO-PLP", "PLP",
    "Bumping", "범핑", "패키징", "후공정", "SiP", "Test", "테스트",
    "PMIC", "DDI", "CIS", "EMI", "Shielding", "차폐",
    "전구체", "Precursor", "ALD", "CVD", "고순도",
    "재배선", "RDL", "Fan-out", "fan out", "wafer level", "panel level",
    "micro bump", "양산", "수율", "공정", "기판", "리드탭",
    "반도체", "소자", "웨이퍼", "방열",
]

FIELD_ALIASES: dict[str, list[str]] = {
    "application_no": ["출원번호", "applicationNumber", "application_number", "application no", "app_no"],
    "publication_no": ["공개번호", "publicationNumber", "publication_number", "openNumber", "open_number"],
    "registration_no": ["등록번호", "registerNumber", "register_number", "registration_number", "reg_no"],
    "title": ["발명의명칭", "발명의 명칭", "inventionTitle", "invention_title", "title", "patent title"],
    "abstract": ["초록", "요약", "abstract", "astrtCont", "summary"],
    "claims": ["청구항", "청구범위", "claim", "claims"],
    "applicant": ["출원인", "출원인명", "권리자", "applicantName", "applicant_name", "applicant", "assignee", "owner"],
    "inventor": ["발명자", "inventor", "inventorName", "inventor_name"],
    "ipc": ["IPC", "ipc", "ipcNumber", "ipc_number", "IPC분류", "국제특허분류", "ipc code"],
    "cpc": ["CPC", "cpc", "CPC분류", "cpc code"],
    "application_date": ["출원일", "출원일자", "applicationDate", "application_date", "filing date"],
    "publication_date": ["공개일", "공개일자", "openDate", "open_date", "publicationDate", "publication_date"],
    "registration_date": ["등록일", "등록일자", "registerDate", "register_date", "registration_date"],
    "status": ["등록상태", "상태", "법적상태", "registerStatus", "register_status", "status", "legal status"],
    "url": ["url", "link", "링크", "kipris url", "detail url"],
    "drawing": ["drawing", "bigDrawing", "big_drawing", "대표도", "도면"],
}

DATE_RE = re.compile(r"(19\d{2}|20\d{2})")
IPC_PREFIX_RE = re.compile(r"\b([A-HY]\d{2}[A-Z])(?:\s*\d{1,4}\s*/\s*\d{1,6})?", re.IGNORECASE)


def _norm_key(value: Any) -> str:
    text = clean_text(value).lower()
    return re.sub(r"[\s_\-()/\\\[\]{}.:：]+", "", text)


ALIAS_TO_FIELD: dict[str, str] = {}
for field, aliases in FIELD_ALIASES.items():
    for alias in aliases:
        ALIAS_TO_FIELD[_norm_key(alias)] = field


def _canonical_field(header: Any) -> str | None:
    key = _norm_key(header)
    if not key:
        return None

    if key in ALIAS_TO_FIELD:
        return ALIAS_TO_FIELD[key]

    for alias_key, field in ALIAS_TO_FIELD.items():
        if alias_key and (alias_key in key or key in alias_key):
            return field

    return None


def _identity_terms(company: dict[str, Any], company_dir: str) -> list[str]:
    terms: list[str] = []

    for key in ["corp_name", "name", "corp_name_en", "company_name", "slug", "stock_code"]:
        value = clean_text(company.get(key))
        if value and value not in terms:
            terms.append(value)

    aliases = company.get("aliases") or []
    if isinstance(aliases, str):
        aliases = [aliases]

    for value in aliases:
        value = clean_text(value)
        if value and value not in terms:
            terms.append(value)

    if company_dir and company_dir not in terms:
        terms.append(company_dir)

    if any("네패스" in term for term in terms):
        for extra in ["네패스", "주식회사 네패스", "(주)네패스", "네패스라웨", "네패스아크", "네패스 하임", "NEPES"]:
            if extra not in terms:
                terms.append(extra)

    return terms


def _tech_terms(company: dict[str, Any]) -> list[str]:
    terms: list[str] = []

    for key in ["core_keywords", "tech_keywords", "products", "keywords"]:
        values = company.get(key) or []
        if isinstance(values, str):
            values = [values]

        for value in values:
            value = clean_text(value)
            if value and value not in terms:
                terms.append(value)

    for value in TECH_BASE_TERMS:
        if value not in terms:
            terms.append(value)

    return terms


def _parse_year(value: Any) -> int | None:
    text = clean_text(value)
    match = DATE_RE.search(text)
    if not match:
        return None

    year = int(match.group(1))
    if 1980 <= year <= datetime.now().year + 1:
        return year
    return None


def _split_ipc_prefixes(value: Any) -> list[str]:
    """
    KIPRIS IPC/CPC 원문 예시:
    - H10W 22/00 | H10B 40/00
    - H01L 21/00; H01L 23/00
    - G01R 31/28 | G06F 11/34

    기존 문제:
    split("/")를 먼저 해버리면 H10W, 22, 40, 49처럼 숫자 토큰이 IPC로 오인됨.

    개선:
    IPC/CPC prefix 형식인 [A-HY][0-9][0-9][A-Z]만 추출한다.
    """
    text = clean_text(value).upper()
    if not text:
        return []

    prefixes: list[str] = []

    for match in IPC_PREFIX_RE.finditer(text):
        prefix = match.group(1).upper().strip()
        if prefix and prefix not in prefixes:
            prefixes.append(prefix)

    return prefixes


def _contains_company(text: str, terms: list[str]) -> bool:
    low = clean_text(text).lower()
    return any(clean_text(term).lower() in low for term in terms if clean_text(term))


def _keyword_hits(text: str, terms: list[str]) -> list[str]:
    low = clean_text(text).lower()
    hits: list[str] = []

    for term in terms:
        t = clean_text(term)
        if t and t.lower() in low and t not in hits:
            hits.append(t)

    return hits


def _read_csv_rows(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []

    for enc in ["utf-8-sig", "utf-8", "cp949", "euc-kr"]:
        try:
            with path.open("r", encoding=enc, errors="ignore", newline="") as f:
                sample = f.read(4096)
                f.seek(0)

                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
                    delimiter = dialect.delimiter
                except Exception:
                    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","

                reader = csv.DictReader(f, delimiter=delimiter)
                for row in reader:
                    if any(clean_text(v) for v in row.values()):
                        rows.append(dict(row))

            if rows:
                return rows

        except Exception:
            continue

    return rows


def _read_excel_rows(path: Path) -> list[dict[str, Any]]:
    if load_workbook is None:
        return []

    rows: list[dict[str, Any]] = []

    try:
        wb = load_workbook(path, data_only=True, read_only=True)

        for ws in wb.worksheets:
            raw_rows = list(ws.iter_rows(values_only=True))
            if not raw_rows:
                continue

            header_idx = 0
            for i, row in enumerate(raw_rows[:10]):
                hits = sum(1 for cell in row if _canonical_field(cell))
                if hits >= 2:
                    header_idx = i
                    break

            headers = [clean_text(v) for v in raw_rows[header_idx]]

            for row in raw_rows[header_idx + 1:]:
                item = {}
                for i, header in enumerate(headers):
                    if not header:
                        continue
                    item[header] = row[i] if i < len(row) else ""

                if any(clean_text(v) for v in item.values()):
                    rows.append(item)

    except Exception:
        return []

    return rows


def _read_json_rows(path: Path) -> list[dict[str, Any]]:
    try:
        data = json.loads(path.read_text(encoding="utf-8", errors="ignore"))
    except Exception:
        return []

    if isinstance(data, list):
        return [x for x in data if isinstance(x, dict)]

    if isinstance(data, dict):
        for key in ["items", "item", "patents", "results", "rows", "data"]:
            value = data.get(key)
            if isinstance(value, list):
                return [x for x in value if isinstance(x, dict)]
        return [data]

    return []


def _read_text_rows(path: Path) -> list[dict[str, Any]]:
    text = ""

    for enc in ["utf-8-sig", "utf-8", "cp949", "euc-kr"]:
        try:
            text = path.read_text(encoding=enc, errors="ignore")
            break
        except Exception:
            continue

    if not text:
        return []

    rows: list[dict[str, Any]] = []
    blocks = re.split(r"\n\s*\n|={3,}|-{3,}", text)

    for block in blocks:
        item: dict[str, Any] = {}

        for line in block.splitlines():
            if ":" in line:
                key, value = line.split(":", 1)
            elif "：" in line:
                key, value = line.split("：", 1)
            else:
                continue

            field = _canonical_field(key)
            if field:
                item[field] = clean_text(value)

        if item:
            rows.append(item)

    if rows:
        return rows

    for i, line in enumerate(text.splitlines(), 1):
        line = clean_text(line)
        if len(line) < 15:
            continue

        if any(k in line for k in ["특허", "출원", "등록", "청구항", "IPC", "CPC", "반도체", "패키징"]):
            rows.append({"title": line[:160], "abstract": line, "source_row": i})

    return rows


def _is_usable_patent_source(path: Path) -> bool:
    if not path.is_file() or path.suffix.lower() not in PATENT_EXTS:
        return False
    name = path.name.lower()
    if path.name.startswith("~$"):
        return False
    if any(skip in name for skip in GENERATED_PATENT_SKIP_KEYWORDS):
        return False
    return any(key.lower() in name for key in PATENT_FILE_KEYWORDS)


def _find_patent_files(company_dir: str) -> list[Path]:
    company_tech_dir = company_agent_dir(company_dir, "tech", create=True)

    source_dirs = [
        company_tech_dir / "source",
        tech_source_dir(company_dir),
        field_common_dir("data") / "patents" / company_dir,
        field_common_dir("data") / "patent" / company_dir,
        field_common_dir("data") / "kipris" / company_dir,
    ]

    source_explicit = [
        company_tech_dir / "source" / f"kipris_{company_dir}_patents.csv",
        company_tech_dir / "source" / f"{company_dir}_kipris_patents.csv",
        company_tech_dir / "source" / f"{company_dir}_kipris_bibliographic.csv",
    ]
    files: list[Path] = [p for p in source_explicit if p.exists() and p.is_file()]

    for base in source_dirs:
        if not base.exists():
            continue
        for path in base.rglob("*"):
            if _is_usable_patent_source(path):
                files.append(path)

    # If no manual/free source export exists, use generated normalized aliases.
    # Do not mix both, because that inflates source counts and duplicate records.
    if not files:
        generated_explicit = [
            company_tech_dir / f"{company_dir}_kipris_bibliographic_normalized.csv",
            company_tech_dir / f"{company_dir}_kipris_patents_normalized.csv",
            company_tech_dir / f"{company_dir}_tech_patent_normalized.csv",
        ]
        files.extend([p for p in generated_explicit if p.exists() and p.is_file()])
        for path in company_tech_dir.glob("*kipris*normalized*.csv"):
            if _is_usable_patent_source(path):
                files.append(path)
        for path in company_tech_dir.glob("*patent*normalized*.csv"):
            if _is_usable_patent_source(path):
                files.append(path)

    # Last-resort fallback: only scan source folders, not generated output root.
    if not files:
        for base in [company_tech_dir / "source", tech_source_dir(company_dir)]:
            if not base.exists():
                continue
            for path in base.rglob("*"):
                if path.is_file() and path.suffix.lower() in {".csv", ".xlsx", ".txt", ".json"}:
                    files.append(path)

    return sorted(set(files))


def _canonicalize_row(row: dict[str, Any]) -> dict[str, Any]:
    out: dict[str, Any] = {}

    for key, value in row.items():
        field = _canonical_field(key)
        if field:
            out[field] = value

    for field in FIELD_ALIASES:
        if field in row and field not in out:
            out[field] = row[field]

    return out


def _normalize_record(
    row: dict[str, Any],
    source_file: Path,
    row_no: int,
    company_terms: list[str],
    tech_terms: list[str],
) -> dict[str, Any]:
    c = _canonicalize_row(row)

    title = clean_text(c.get("title"))
    applicant = clean_text(c.get("applicant"))
    abstract = clean_text(c.get("abstract"))
    claims = clean_text(c.get("claims"))
    status = clean_text(c.get("status"))

    application_no = clean_text(c.get("application_no"))
    publication_no = clean_text(c.get("publication_no"))
    registration_no = clean_text(c.get("registration_no"))

    application_date = clean_text(c.get("application_date"))
    publication_date = clean_text(c.get("publication_date"))
    registration_date = clean_text(c.get("registration_date"))

    ipc_codes = _split_ipc_prefixes(c.get("ipc"))
    cpc_codes = _split_ipc_prefixes(c.get("cpc"))

    joined = " ".join([title, applicant, abstract, claims, status, application_no, registration_no])
    applicant_matched = _contains_company(joined, company_terms)
    tech_hits = _keyword_hits(" ".join([title, abstract, claims, " ".join(ipc_codes), " ".join(cpc_codes)]), tech_terms)

    app_year = _parse_year(application_date)
    pub_year = _parse_year(publication_date)
    reg_year = _parse_year(registration_date)

    is_registered = bool(registration_no) or any(k in status for k in ["등록", "존속"])
    is_expired_or_dead = any(k in status for k in ["소멸", "거절", "취하", "포기", "무효", "말소"])

    record_key = "|".join([application_no, registration_no, title[:80], applicant[:80]])

    return {
        "record_key": record_key,
        "title": title,
        "applicant": applicant,
        "inventor": clean_text(c.get("inventor")),
        "application_no": application_no,
        "publication_no": publication_no,
        "registration_no": registration_no,
        "application_date": application_date,
        "publication_date": publication_date,
        "registration_date": registration_date,
        "application_year": app_year,
        "publication_year": pub_year,
        "registration_year": reg_year,
        "status": status,
        "is_registered": is_registered,
        "is_expired_or_dead": is_expired_or_dead,
        "ipc_codes": ipc_codes,
        "cpc_codes": cpc_codes,
        "abstract": abstract,
        "claims": claims,
        "url": clean_text(c.get("url")),
        "drawing": clean_text(c.get("drawing")),
        "applicant_matched": applicant_matched,
        "tech_keyword_hits": tech_hits,
        "source_file": str(source_file),
        "source_row": row_no,
    }


def _dedupe(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    seen: set[str] = set()
    out: list[dict[str, Any]] = []

    for r in records:
        key = clean_text(r.get("record_key"))
        if not key or key == "|||":
            key = "|".join([
                clean_text(r.get("title")),
                clean_text(r.get("applicant")),
                clean_text(r.get("application_date")),
            ])

        if key in seen:
            continue

        seen.add(key)
        out.append(r)

    return out


def _rank_records(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    current_year = datetime.now().year
    ranked: list[dict[str, Any]] = []

    for r in records:
        score = 0.0

        if r.get("applicant_matched"):
            score += 20
        if r.get("is_registered"):
            score += 10
        if not r.get("is_expired_or_dead"):
            score += 5

        year = r.get("application_year") or r.get("registration_year")
        if isinstance(year, int):
            if year >= current_year - 5:
                score += 8
            elif year >= current_year - 10:
                score += 4

        score += min(len(r.get("tech_keyword_hits") or []) * 3, 18)
        score += min(len(r.get("ipc_codes") or []) * 1.5, 6)

        if clean_text(r.get("abstract")):
            score += 3
        if clean_text(r.get("claims")):
            score += 4

        item = dict(r)
        item["patent_rank_score"] = round(score, 2)
        ranked.append(item)

    return sorted(ranked, key=lambda x: x.get("patent_rank_score", 0), reverse=True)


def _metric(name: str, value: Any, unit: str, formula: str, detail: Any = None) -> dict[str, Any]:
    return {
        "metric_name": name,
        "value": value,
        "unit": unit,
        "formula": formula,
        "detail": detail,
    }


def _build_metrics(records: list[dict[str, Any]], all_records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    current_year = datetime.now().year

    company_matched = [r for r in records if r.get("applicant_matched")]
    base_records = company_matched if company_matched else records

    registered = [r for r in base_records if r.get("is_registered")]
    active = [r for r in registered if not r.get("is_expired_or_dead")]

    recent = []
    for r in base_records:
        year = r.get("application_year") or r.get("registration_year")
        if isinstance(year, int) and year >= current_year - 5:
            recent.append(r)

    ipc_prefixes: list[str] = []
    h01l_count = 0
    semiconductor_ipc_count = 0

    for r in base_records:
        prefixes = set((r.get("ipc_codes") or []) + (r.get("cpc_codes") or []))

        for prefix in sorted(prefixes):
            if prefix and prefix not in ipc_prefixes:
                ipc_prefixes.append(prefix)

        if "H01L" in prefixes:
            h01l_count += 1

        if prefixes & SEMICONDUCTOR_IPC_PREFIXES:
            semiconductor_ipc_count += 1

    keyword_counter: Counter[str] = Counter()
    for r in base_records:
        keyword_counter.update(r.get("tech_keyword_hits") or [])

    years = [
        r.get("application_year") or r.get("registration_year")
        for r in base_records
        if isinstance(r.get("application_year") or r.get("registration_year"), int)
    ]
    year_span = f"{min(years)}~{max(years)}" if years else "연도 미확인"

    return [
        _metric("특허 원천 파일 수", len(set(clean_text(r.get("source_file")) for r in all_records)), "개", "KIPRIS/특허 파일명으로 인식된 원천 파일 수"),
        _metric("정규화 특허 레코드 수", len(records), "건", "KIPRIS CSV/XLSX/TXT/JSON을 표준 특허 레코드로 변환 후 중복 제거한 수"),
        _metric("출원인/권리자 회사 매칭 특허 수", len(company_matched), "건", "출원인·권리자·본문에서 회사명/별칭이 확인된 특허 수"),
        _metric("등록 특허 수", len(registered), "건", "등록번호가 있거나 등록상태가 확인된 회사 매칭 특허 수"),
        _metric("존속 가능 특허 수", len(active), "건", "등록 특허 중 소멸·취하·거절·포기·무효로 분류되지 않은 특허 수"),
        _metric("최근 5년 특허 수", len(recent), "건", f"출원/등록 연도가 {current_year - 5}년 이후인 회사 매칭 특허 수"),
        _metric("IPC/CPC 기술분류 수", len(ipc_prefixes), "개", "회사 매칭 특허의 IPC/CPC prefix 고유 개수", ipc_prefixes[:30]),
        _metric("H01L 반도체 IPC 특허 수", h01l_count, "건", "IPC/CPC prefix가 H01L인 반도체 장치·공정 관련 특허 수"),
        _metric("반도체 핵심 IPC 특허 수", semiconductor_ipc_count, "건", "H01L/H10B/H10W/G01R/G06F/C23C 등 반도체 관련 IPC/CPC prefix가 포함된 특허 수"),
        _metric("특허-기술 키워드 매칭 수", sum(keyword_counter.values()), "회", "특허 제목·초록·청구항 내 기술 키워드 출현 횟수", dict(keyword_counter.most_common(15))),
        _metric("특허 포트폴리오 연도 범위", year_span, "", "회사 매칭 특허의 출원/등록 연도 범위"),
    ]


def _patent_to_text(record: dict[str, Any]) -> str:
    title = clean_text(record.get("title")) or "특허명 미확인"
    applicant = clean_text(record.get("applicant")) or "출원인 미확인"
    status = clean_text(record.get("status")) or "상태 미확인"
    app_no = clean_text(record.get("application_no"))
    reg_no = clean_text(record.get("registration_no"))
    year = record.get("application_year") or record.get("registration_year") or "연도 미확인"
    ipc = ", ".join((record.get("ipc_codes") or [])[:6]) or "IPC 미확인"
    keywords = ", ".join((record.get("tech_keyword_hits") or [])[:8]) or "기술 키워드 미확인"

    nums = []
    if app_no:
        nums.append(f"출원번호 {app_no}")
    if reg_no:
        nums.append(f"등록번호 {reg_no}")

    abstract = clean_text(record.get("abstract"))[:220]
    text = f"{title} | {applicant} | {status} | {year} | {'; '.join(nums)} | IPC/CPC: {ipc} | 기술매칭: {keywords}"
    if abstract:
        text += f" | 초록: {abstract}"
    return text


def _build_docs(
    company: dict[str, Any],
    company_dir: str,
    metrics: list[dict[str, Any]],
    representatives: list[dict[str, Any]],
    source_files: list[Path],
) -> list[dict[str, Any]]:
    company_name = clean_text(company.get("corp_name") or company.get("name") or company_dir)

    metric_lines = [f"- {m['metric_name']}: {m['value']}{m['unit']}" for m in metrics]
    rep_lines = [f"{idx}. {_patent_to_text(r)}" for idx, r in enumerate(representatives[:10], 1)]

    summary = "\n".join([
        f"{company_name} KIPRIS/특허 정량화 요약",
        "본 특허 근거는 KIPRIS API/CSV/엑셀/TXT 원천 파일을 정규화해 산출했다.",
        "",
        "[정량 지표]",
        *metric_lines,
        "",
        "[대표 특허]",
        *(rep_lines or ["대표 특허 없음"]),
    ])

    docs: list[dict[str, Any]] = [
        {
            "doc_id": "tech.patent.summary",
            "source_type": "patent",
            "title": f"{company_name} KIPRIS/특허 정량화 요약",
            "path": ", ".join(str(p) for p in source_files[:5]),
            "url": "",
            "text": summary,
            "relevance_score": 30,
            "relevance_reasons": ["KIPRIS", "특허 정량화", "출원인/권리자 매칭", "IPC/CPC 분석"],
        }
    ]

    for idx, r in enumerate(representatives[:20], 1):
        docs.append(
            {
                "doc_id": f"tech.patent.{idx:03d}",
                "source_type": "patent",
                "title": clean_text(r.get("title"))[:180] or f"대표 특허 {idx}",
                "path": clean_text(r.get("source_file")),
                "url": clean_text(r.get("url")),
                "text": _patent_to_text(r),
                "relevance_score": 20 + float(r.get("patent_rank_score") or 0),
                "relevance_reasons": [
                    "대표 특허 레코드",
                    "회사 매칭" if r.get("applicant_matched") else "회사 매칭 약함",
                    f"기술 키워드 {len(r.get('tech_keyword_hits') or [])}개",
                ],
            }
        )

    return docs


def harvest_patent_evidence(company: dict[str, Any], company_dir: str) -> dict[str, Any]:
    company_terms = _identity_terms(company, company_dir)
    tech_terms = _tech_terms(company)
    source_files = _find_patent_files(company_dir)

    raw_records: list[dict[str, Any]] = []

    for path in source_files:
        suffix = path.suffix.lower()

        if suffix in {".csv", ".tsv"}:
            rows = _read_csv_rows(path)
        elif suffix in {".xlsx", ".xlsm"}:
            rows = _read_excel_rows(path)
        elif suffix in {".txt", ".md"}:
            rows = _read_text_rows(path)
        elif suffix == ".json":
            rows = _read_json_rows(path)
        else:
            rows = []

        for idx, row in enumerate(rows, 1):
            record = _normalize_record(row, path, idx, company_terms, tech_terms)
            has_content = any(
                clean_text(record.get(k))
                for k in ["title", "abstract", "claims", "application_no", "registration_no", "applicant"]
            )
            if has_content:
                raw_records.append(record)

    normalized = _dedupe(raw_records)
    ranked = _rank_records(normalized)

    company_matched = [r for r in ranked if r.get("applicant_matched")]
    effective = company_matched if company_matched else ranked

    metrics = _build_metrics(effective, normalized)

    kipris_tech_ml_features = _load_kipris_tech_ml_features(company, locals().get('company_dir'))

    if kipris_tech_ml_features:

        metrics = _merge_kipris_tech_ml_metrics(metrics, kipris_tech_ml_features)
    representatives = effective[:20]
    docs = _build_docs(company, company_dir, metrics, representatives, source_files) if normalized else []

    quality_flags: list[str] = []
    if not source_files:
        quality_flags.append("KIPRIS/특허 원천 파일이 없습니다.")
    if source_files and not normalized:
        quality_flags.append("특허 원천 파일은 있으나 정규화 가능한 레코드가 없습니다.")
    if normalized and not company_matched:
        quality_flags.append("특허 레코드는 있으나 출원인/권리자 회사명 매칭이 약합니다. applicant 검색어를 확인하세요.")

    return {
        "version": "v14_kipris_ipc_refined",
        "company_dir": company_dir,
        "company_name": clean_text(company.get("corp_name") or company.get("name") or company_dir),
        "source_files": [str(p) for p in source_files],
        "raw_record_count": len(raw_records),
        "normalized_record_count": len(normalized),
        "company_matched_record_count": len(company_matched),
        "effective_record_count": len(effective),
        "records": ranked[:500],
        "representative_patents": representatives,
        "metrics": metrics,
        "kipris_tech_ml_features": kipris_tech_ml_features,
        "docs": docs,
        "quality_flags": quality_flags,
    }


def _metric_line(metric: dict[str, Any]) -> str:
    name = clean_text(metric.get("metric_name"))
    value = clean_text(metric.get("value"))
    unit = clean_text(metric.get("unit"))
    detail = metric.get("detail")

    suffix = ""
    if isinstance(detail, dict) and detail:
        suffix = " — " + ", ".join(f"{k} {v}회" for k, v in list(detail.items())[:10])
    elif isinstance(detail, list) and detail:
        suffix = " — " + ", ".join(clean_text(x) for x in detail[:12])

    return f"- {name}: {value}{unit}{suffix}"


def write_patent_outputs(
    company_dir: str,
    patent_harvest: dict[str, Any] | None,
    output_dir: str | Path = OUTPUT_DIR,
) -> dict[str, str]:
    patent_harvest = patent_harvest or {}
    out_dir = ensure_dir(Path(output_dir))

    json_path = out_dir / f"{company_dir}_tech_patent_evidence.json"
    md_path = out_dir / f"{company_dir}_tech_patent_evidence.md"
    csv_path = out_dir / f"{company_dir}_tech_patent_normalized.csv"

    json_path.write_text(json.dumps(patent_harvest, ensure_ascii=False, indent=2), encoding="utf-8")

    company_name = patent_harvest.get("company_name") or company_dir

    lines = [
        f"# {company_name} KIPRIS/특허 정량화 결과",
        "",
        "## 1. KIPRIS/특허 정량화 요약",
        f"- 원천 파일 수: {len(patent_harvest.get('source_files') or [])}",
        f"- 원천 특허 레코드 수: {patent_harvest.get('raw_record_count', 0)}",
        f"- 정규화 특허 레코드 수: {patent_harvest.get('normalized_record_count', 0)}",
        f"- 회사 매칭 특허 레코드 수: {patent_harvest.get('company_matched_record_count', 0)}",
        "",
        "## 2. 특허 정량 지표",
    ]

    for metric in patent_harvest.get("metrics") or []:
        lines.append(_metric_line(metric))

    lines += ["", "## 3. 대표 특허"]

    representatives = patent_harvest.get("representative_patents") or []
    if representatives:
        for idx, record in enumerate(representatives[:20], 1):
            lines.append(f"{idx}. {_patent_to_text(record)}")
    else:
        lines.append("- 대표 특허 없음")

    lines += ["", "## 4. 원천 파일"]
    for file in patent_harvest.get("source_files") or []:
        lines.append(f"- {file}")

    lines += ["", "## 5. 품질 플래그"]
    flags = patent_harvest.get("quality_flags") or []
    if flags:
        for flag in flags:
            lines.append(f"- {flag}")
    else:
        lines.append("- 특이사항 없음")

    md_path.write_text("\n".join(lines), encoding="utf-8")

    records = patent_harvest.get("records") or []
    fieldnames = [
        "record_key", "title", "applicant", "application_no", "publication_no", "registration_no",
        "status", "application_year", "publication_year", "registration_year",
        "ipc_codes", "cpc_codes", "tech_keyword_hits", "applicant_matched",
        "is_registered", "is_expired_or_dead", "patent_rank_score", "source_file", "source_row",
    ]

    with csv_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()

        for record in records:
            row = {k: record.get(k, "") for k in fieldnames}
            for key in ["ipc_codes", "cpc_codes", "tech_keyword_hits"]:
                if isinstance(row.get(key), list):
                    row[key] = "; ".join(clean_text(x) for x in row[key])
            writer.writerow(row)

    return {"json": str(json_path), "md": str(md_path), "csv": str(csv_path)}


# KIPRIS_TECH_ML_RUNTIME_PATCH_V2_START
def _find_kipris_tech_ml_feature_file(company: str, company_dir: str | None = None) -> Path | None:
    """Find prebuilt KIPRIS Tech ML feature JSON.

    Expected example:
    data/반도체/네패스/tech/nepes_kipris_tech_ml_features.json
    """
    root = Path.cwd()
    data_dir = root / "data"

    candidates: list[Path] = []

    if company_dir:
        candidates.extend(data_dir.glob(f"*/*/tech/{company_dir}_kipris_tech_ml_features.json"))
        candidates.extend(data_dir.glob(f"*/{company}/tech/{company_dir}_kipris_tech_ml_features.json"))

    if company:
        candidates.extend(data_dir.glob(f"*/{company}/tech/*_kipris_tech_ml_features.json"))

    candidates.extend(data_dir.glob("*/*/tech/*_kipris_tech_ml_features.json"))

    seen: set[str] = set()
    unique_candidates: list[Path] = []
    for path in candidates:
        key = str(path.resolve())
        if key not in seen:
            seen.add(key)
            unique_candidates.append(path)

    for path in unique_candidates:
        if path.exists() and path.is_file():
            return path

    return None


def _load_kipris_tech_ml_features(company: str, company_dir: str | None = None) -> dict[str, Any] | None:
    path = _find_kipris_tech_ml_feature_file(company, company_dir)
    if not path:
        return None

    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return None

    if not isinstance(data, dict):
        return None

    data["_source_path"] = str(path)
    return data


def _pct_from_rate(value: Any) -> float | None:
    try:
        if value is None:
            return None
        return round(float(value) * 100.0, 2)
    except Exception:
        return None


def _num_or_none(value: Any, digits: int = 2) -> float | None:
    try:
        if value is None:
            return None
        return round(float(value), digits)
    except Exception:
        return None


def _merge_kipris_tech_ml_metrics(metrics: list[dict[str, Any]], features: dict[str, Any] | None) -> list[dict[str, Any]]:
    """Merge normalized KIPRIS Tech ML features into patent metrics.

    This does not treat patents as direct commercialization evidence.
    It only adds conservative IP quality / legal stability / technology-fit signals.
    """
    if not features:
        return metrics

    metrics = list(metrics or [])
    existing_names = {str(m.get("metric_name") or "") for m in metrics if isinstance(m, dict)}

    core_counts = features.get("core_counts") or {}
    core_rates = features.get("core_rates") or {}
    scores = features.get("scores") or {}
    adjustment = features.get("tech_to_value_bridge_adjustment") or {}
    source_path = features.get("_source_path") or features.get("source_path") or "KIPRIS Tech ML features"

    def add_metric(metric_name: str, value: Any, unit: str, detail: str) -> None:
        if metric_name in existing_names:
            return
        if value is None:
            return
        metrics.append(
            {
                "metric_name": metric_name,
                "value": value,
                "unit": unit,
                "detail": detail,
                "source_type": "kipris_tech_ml_features",
                "source": source_path,
            }
        )
        existing_names.add(metric_name)

    add_metric(
        "KIPRIS 총 특허 수",
        _num_or_none(core_counts.get("total_patents"), 0),
        "건",
        "정규화된 KIPRIS 원천 CSV 기준 전체 특허 레코드 수입니다.",
    )
    add_metric(
        "KIPRIS 등록 특허 수(추정)",
        _num_or_none(core_counts.get("registered_patents_estimated"), 0),
        "건",
        "register_status/register_number/register_date 기반 등록 특허 수 추정치입니다.",
    )
    add_metric(
        "KIPRIS 존속 가능 특허 수(추정)",
        _num_or_none(core_counts.get("alive_patents_estimated"), 0),
        "건",
        "등록 상태와 권리상태 추정 로직을 반영한 존속 가능 특허 수입니다.",
    )
    add_metric(
        "KIPRIS 등록률(추정)",
        _pct_from_rate(core_rates.get("registration_rate_estimated")),
        "%",
        "등록 특허 수를 전체 특허 수로 나눈 보수적 등록률 추정치입니다.",
    )
    add_metric(
        "KIPRIS 존속률(전체 대비 추정)",
        _pct_from_rate(core_rates.get("alive_rate_total_estimated")),
        "%",
        "존속 가능 특허 수를 전체 특허 수로 나눈 보수적 존속률 추정치입니다.",
    )
    add_metric(
        "KIPRIS 존속률(등록특허 중 추정)",
        _pct_from_rate(core_rates.get("alive_rate_among_registered_estimated")),
        "%",
        "존속 가능 특허 수를 등록 특허 수로 나눈 권리 안정성 지표입니다.",
    )
    add_metric(
        "KIPRIS 부정 처분 비중(추정)",
        _pct_from_rate(core_rates.get("negative_disposal_rate_estimated")),
        "%",
        "미등록·소멸·거절·취하 등으로 추정되는 부정 처분성 상태 비중입니다.",
    )
    add_metric(
        "KIPRIS 최근 5년 출원 비중",
        _pct_from_rate(core_rates.get("recent_5y_application_rate")),
        "%",
        "최근 5년 출원 특허 비중으로 포트폴리오의 최신성을 나타냅니다.",
    )
    add_metric(
        "KIPRIS 반도체 관련 IPC 특허 수",
        _num_or_none(core_counts.get("semiconductor_related_ipc_patents"), 0),
        "건",
        "반도체 관련 IPC/CPC 분류와 연결되는 특허 수입니다.",
    )
    add_metric(
        "KIPRIS 반도체 관련 IPC 비중",
        _pct_from_rate(core_rates.get("semiconductor_related_ipc_rate")),
        "%",
        "전체 특허 중 반도체 관련 IPC/CPC 분류 특허 비중입니다.",
    )
    add_metric(
        "KIPRIS H01L 핵심 IPC 특허 수",
        _num_or_none(core_counts.get("h01l_core_patents"), 0),
        "건",
        "H01L 반도체 소자·패키징 관련 핵심 IPC 특허 수입니다.",
    )
    add_metric(
        "KIPRIS 법적 안정성 점수",
        _num_or_none(scores.get("legal_stability_score_estimated"), 2),
        "점",
        "등록률·존속률·부정 처분 비중을 반영한 권리 안정성 점수입니다.",
    )
    add_metric(
        "KIPRIS 포트폴리오 모멘텀 점수",
        _num_or_none(scores.get("portfolio_momentum_score"), 2),
        "점",
        "최근 출원 비중과 포트폴리오 최신성을 반영한 점수입니다.",
    )
    add_metric(
        "KIPRIS 기술 적합도 점수",
        _num_or_none(scores.get("ip_technology_fit_score"), 2),
        "점",
        "반도체 IPC 비중, H01L 특허, 기술 키워드 매칭을 반영한 점수입니다.",
    )
    add_metric(
        "KIPRIS Tech ML 종합 점수",
        _num_or_none(scores.get("kipris_tech_ml_score"), 2),
        "점",
        "권리 안정성, 포트폴리오 모멘텀, 기술 적합도를 종합한 KIPRIS 기반 Tech ML 점수입니다.",
    )
    add_metric(
        "KIPRIS Tech-to-Value Bridge 보정점",
        _num_or_none(adjustment.get("bridge_adjustment_points"), 2),
        "점",
        "직접 사업화 근거가 아니라 IP 품질 보조 신호로만 사용하는 보수적 Bridge 보정점입니다.",
    )

    return metrics
# KIPRIS_TECH_ML_RUNTIME_PATCH_V2_END

