from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None

from .config import DEFAULT_TEMPLATE_PATH, OUTPUT_DIR
from .utils import clean_text, ensure_dir

NUM_RE = re.compile(r"(?P<value>[-+]?\d{1,3}(?:,\d{3})*(?:\.\d+)?|[-+]?\d+(?:\.\d+)?)\s*(?P<unit>%|배|건|개|개사|곳|억원|백만원|년|개월|회|종|명|라인|단|nm|㎚|um|μm|㎛)?")


def _vals(row: tuple[Any, ...]) -> list[str]:
    return [clean_text(v) for v in row]


def _idx(headers: list[str], keys: list[str], default: int) -> int:
    for i, h in enumerate(headers):
        hl = h.lower()
        if any(k.lower() in hl for k in keys):
            return i
    return default


def load_formula_rules(template_path: str | Path = DEFAULT_TEMPLATE_PATH) -> list[dict[str, Any]]:
    path = Path(template_path)
    if not path.exists() or load_workbook is None:
        return []
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception:
        return []
    rules: list[dict[str, Any]] = []
    targets = [s for s in wb.sheetnames if "수식" in s or "정량" in s or "formula" in s.lower()]
    for sname in targets:
        ws = wb[sname]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header_i, headers = 0, _vals(rows[0])
        for i, row in enumerate(rows[:10]):
            v = _vals(row)
            joined = " ".join(v).lower()
            if any(k in joined for k in ["지표", "metric", "수식", "formula", "산출"]):
                header_i, headers = i, v
                break
        ci = _idx(headers, ["대분류", "category", "분류", "축"], 0)
        mi = _idx(headers, ["지표", "metric", "항목", "name"], 1)
        fi = _idx(headers, ["수식", "formula", "산출", "계산"], 2)
        ii = _idx(headers, ["의미", "해석", "interpret"], 3)
        si = _idx(headers, ["출처", "source", "근거", "data"], 4)
        for row in rows[header_i + 1:]:
            v = _vals(row)
            if not any(v):
                continue
            metric = v[mi] if mi < len(v) else ""
            formula = v[fi] if fi < len(v) else ""
            if not metric and not formula:
                continue
            rules.append({
                "category": v[ci] if ci < len(v) and v[ci] else "공통",
                "metric_name": metric or formula[:40],
                "formula": formula,
                "interpretation": v[ii] if ii < len(v) else "",
                "source_hint": v[si] if si < len(v) else "",
                "sheet": sname,
            })
    return rules


def _keywords(rule: dict[str, Any]) -> list[str]:
    text = " ".join(clean_text(rule.get(k)) for k in ["category", "metric_name", "formula", "interpretation", "source_hint"])
    toks = re.split(r"[^가-힣A-Za-z0-9μ㎛%]+", text)
    out: list[str] = []
    for t in toks:
        if len(t) >= 2 and t.lower() not in {"and", "or", "the", "ratio", "rate"} and t not in out:
            out.append(t)
    return out[:8]


def extract_quantitative_signals(text: str, rules: list[dict[str, Any]], max_signals: int = 120) -> list[dict[str, Any]]:
    text = clean_text(text)[:40000]
    if not text:
        return []
    sents = [clean_text(x) for x in re.split(r"(?<=[.!?。다])\s+|\n+", text) if len(clean_text(x)) > 10]
    out: list[dict[str, Any]] = []
    for rule in rules:
        kws = _keywords(rule)
        if not kws:
            continue
        hit_count = 0
        for sent in sents:
            low = sent.lower()
            if not any(k.lower() in low for k in kws):
                continue
            nums = []
            for m in NUM_RE.finditer(sent[:600]):
                start = max(0, m.start() - 20)
                nums.append({"label": clean_text(sent[start:m.start()])[-30:], "value": m.group("value"), "unit": clean_text(m.group("unit"))})
                if len(nums) >= 4:
                    break
            if nums:
                item = dict(rule)
                item.update({"numbers": nums, "snippet": sent[:420]})
                out.append(item)
                hit_count += 1
                if hit_count >= 3 or len(out) >= max_signals:
                    break
        if len(out) >= max_signals:
            break
    return out


def write_formula_catalog(company_dir: str, template_path: str | Path = DEFAULT_TEMPLATE_PATH, output_dir: str | Path = OUTPUT_DIR) -> dict[str, str]:
    rules = load_formula_rules(template_path)
    out = ensure_dir(Path(output_dir))
    jp = out / f"{company_dir}_tech_formula_rules.json"
    mp = out / f"{company_dir}_tech_formula_rules.md"
    payload = {"template_path": str(template_path), "formula_rule_count": len(rules), "rules": rules}
    jp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    lines = ["# Tech Agent 수식 정리 코드화", "", f"- rule_count: {len(rules)}"]
    for r in rules:
        lines += ["", f"## {r.get('category')} - {r.get('metric_name')}"]
        if r.get("formula"):
            lines.append(f"- 산식: {r.get('formula')}")
        if r.get("interpretation"):
            lines.append(f"- 해석: {r.get('interpretation')}")
        if r.get("source_hint"):
            lines.append(f"- 추출 출처: {r.get('source_hint')}")
    mp.write_text("\n".join(lines), encoding="utf-8")
    return {"json": str(jp), "md": str(mp)}
