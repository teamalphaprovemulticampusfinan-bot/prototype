from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .utils import clean_text
from .workbook_schema import CANONICAL_SECTION_ORDER, SECTION_FIELD_GUIDE


THIN = Side(style="thin", color="D9D9D9")
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

FILL_HEADER = PatternFill("solid", fgColor="DCE6F1")
FILL_SUB = PatternFill("solid", fgColor="EEF3F8")
FILL_SECTION = PatternFill("solid", fgColor="F7F7F7")

FONT_HEADER = Font(bold=True, size=11)
FONT_TITLE = Font(bold=True, size=13)
FONT_BOLD = Font(bold=True)

ALIGN_TOP = Alignment(vertical="top", wrap_text=True)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_excel_report(
    company: dict | None = None,
    categories: list[dict] | None = None,
    score_payload: dict | None = None,
    output_path=None,
    template_path=None,
    **kwargs,
):
    """
    호환용 래퍼.
    기존 main.py에서 아래 형태로 들어와도 받는다.
    - score_result=
    - out_path=
    - template_path=None
    """
    if company is None:
        company = kwargs.get("company", {})
    if categories is None:
        categories = kwargs.get("categories", [])
    if score_payload is None:
        score_payload = kwargs.get("score_result") or kwargs.get("score_payload") or {}
    if output_path is None:
        output_path = kwargs.get("out_path") or kwargs.get("output_path")
    if template_path is None:
        template_path = kwargs.get("template_path")

    if output_path is None:
        raise ValueError("output_path 또는 out_path가 필요합니다.")

    wb = _load_or_create_workbook(template_path)

    _prepare_base_sheets(wb)
    _write_summary_sheet(wb["요약"], company, score_payload)
    _write_analysis_sheet(wb["기술분석"], company, categories, score_payload)
    _write_quant_sheet(wb["정량지표"], company, categories, score_payload)
    _write_source_sheet(wb["원천문서"], company)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def _load_or_create_workbook(template_path=None):
    if template_path:
        template_path = Path(template_path)
        if template_path.exists():
            return load_workbook(template_path)

    wb = Workbook()
    default_ws = wb.active
    default_ws.title = "요약"
    wb.create_sheet("기술분석")
    wb.create_sheet("정량지표")
    wb.create_sheet("원천문서")
    return wb


def _prepare_base_sheets(wb):
    wanted = ["요약", "기술분석", "정량지표", "원천문서"]
    existing = wb.sheetnames

    for name in wanted:
        if name not in existing:
            wb.create_sheet(name)

    for name in wanted:
        ws = wb[name]
        _clear_sheet(ws)
        _apply_default_layout(ws)


def _clear_sheet(ws):
    if ws.max_row > 1 or ws.max_column > 1 or ws["A1"].value is not None:
        ws.delete_rows(1, ws.max_row)


def _apply_default_layout(ws):
    widths = {
        "A": 18,
        "B": 24,
        "C": 32,
        "D": 60,
        "E": 60,
        "F": 22,
        "G": 18,
        "H": 18,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"


def _set_cell(ws, row, col, value, *, fill=None, font=None, border=True, align=None):
    cell = ws.cell(row=row, column=col, value=value)
    if fill:
        cell.fill = fill
    if font:
        cell.font = font
    if border:
        cell.border = BORDER_ALL
    cell.alignment = align or ALIGN_TOP
    return cell


def _find_section(categories: list[dict], section_name: str) -> dict:
    for row in categories or []:
        if clean_text(row.get("category", "")) == section_name:
            return row
    return {
        "category": section_name,
        "result": "공개 자료 기준 직접 확인이 제한되어 보수적으로 정리함.",
        "items": [],
        "quant_points": [],
        "source_types": [],
    }


def _score_map(score_payload: dict) -> dict[str, dict]:
    result = {}
    for row in score_payload.get("axes", []) or []:
        section = clean_text(row.get("section", ""))
        axis = clean_text(row.get("axis", ""))
        if section:
            result[section] = row
        elif axis:
            result[axis] = row
    return result


def _write_summary_sheet(ws, company: dict, score_payload: dict):
    corp_name = clean_text(company.get("corp_name", "기업"))
    corp_name_en = clean_text(company.get("corp_name_en", ""))
    stock_code = clean_text(company.get("stock_code", ""))

    total_score = score_payload.get("total_score", 0)
    max_score = score_payload.get("max_score", 35)
    strengths = ", ".join(score_payload.get("strengths", []) or []) or "상대적 강점 축 제한적"
    weaknesses = ", ".join(score_payload.get("weaknesses", []) or []) or "상대적 약점 축 제한적"
    judgment = clean_text(score_payload.get("judgment", "보수 검토형"))

    ws.merge_cells("A1:H1")
    _set_cell(
        ws,
        1,
        1,
        f"{corp_name} 기술 분석 요약",
        fill=FILL_HEADER,
        font=FONT_TITLE,
        align=ALIGN_CENTER,
    )

    rows = [
        ("기업명", corp_name),
        ("영문명", corp_name_en or "-"),
        ("종목코드", stock_code or "-"),
        ("총점", f"{total_score}/{max_score}"),
        ("강점", strengths),
        ("약점", weaknesses),
        ("종합 판단", judgment),
    ]

    r = 3
    for label, value in rows:
        _set_cell(ws, r, 1, label, fill=FILL_SUB, font=FONT_BOLD)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
        _set_cell(ws, r, 2, value)
        r += 1

    r += 1
    _set_cell(ws, r, 1, "평가축", fill=FILL_HEADER, font=FONT_HEADER, align=ALIGN_CENTER)
    _set_cell(ws, r, 2, "연결 섹션", fill=FILL_HEADER, font=FONT_HEADER, align=ALIGN_CENTER)
    _set_cell(ws, r, 3, "점수", fill=FILL_HEADER, font=FONT_HEADER, align=ALIGN_CENTER)
    _set_cell(ws, r, 4, "판단 근거", fill=FILL_HEADER, font=FONT_HEADER, align=ALIGN_CENTER)
    _set_cell(ws, r, 5, "정량 신호 수", fill=FILL_HEADER, font=FONT_HEADER, align=ALIGN_CENTER)
    _set_cell(ws, r, 6, "근거 신호 수", fill=FILL_HEADER, font=FONT_HEADER, align=ALIGN_CENTER)

    r += 1
    for row in score_payload.get("axes", []) or []:
        _set_cell(ws, r, 1, clean_text(row.get("axis", "")))
        _set_cell(ws, r, 2, clean_text(row.get("section", "")))
        _set_cell(ws, r, 3, row.get("score", 0), align=ALIGN_CENTER)
        _set_cell(ws, r, 4, clean_text(row.get("reason", "")))
        _set_cell(ws, r, 5, row.get("quant_count", 0), align=ALIGN_CENTER)
        _set_cell(ws, r, 6, row.get("evidence_count", 0), align=ALIGN_CENTER)
        r += 1


def _write_analysis_sheet(ws, company: dict, categories: list[dict], score_payload: dict):
    corp_name = clean_text(company.get("corp_name", "기업"))
    score_lookup = _score_map(score_payload)

    ws.merge_cells("A1:H1")
    _set_cell(
        ws,
        1,
        1,
        f"{corp_name} 기술 분석 상세",
        fill=FILL_HEADER,
        font=FONT_TITLE,
        align=ALIGN_CENTER,
    )

    row = 3
    for idx, section_name in enumerate(CANONICAL_SECTION_ORDER, start=1):
        section = _find_section(categories, section_name)
        linked_score = score_lookup.get(section_name, {})

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        _set_cell(
            ws,
            row,
            1,
            f"{idx}. {section_name}",
            fill=FILL_SECTION,
            font=FONT_HEADER,
        )
        row += 1

        _set_cell(ws, row, 1, "평가축", fill=FILL_SUB, font=FONT_BOLD)
        _set_cell(ws, row, 2, clean_text(linked_score.get("axis", "-")))
        _set_cell(ws, row, 3, "점수", fill=FILL_SUB, font=FONT_BOLD)
        _set_cell(ws, row, 4, linked_score.get("score", "-"), align=ALIGN_CENTER)
        _set_cell(ws, row, 5, "판단 근거", fill=FILL_SUB, font=FONT_BOLD)
        ws.merge_cells(start_row=row, start_column=6, end_row=row, end_column=8)
        _set_cell(ws, row, 6, clean_text(linked_score.get("reason", "-")))
        row += 1

        _set_cell(ws, row, 1, "추출 항목", fill=FILL_SUB, font=FONT_BOLD)
        fields = SECTION_FIELD_GUIDE.get(section_name, [])
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        _set_cell(ws, row, 2, " / ".join(fields) if fields else "-")
        row += 1

        _set_cell(ws, row, 1, "결과", fill=FILL_SUB, font=FONT_BOLD)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        _set_cell(
            ws,
            row,
            2,
            clean_text(section.get("result", "공개 자료 기준 직접 확인이 제한되어 보수적으로 정리함.")),
        )
        row += 1

        _set_cell(ws, row, 1, "항목명", fill=FILL_HEADER, font=FONT_HEADER, align=ALIGN_CENTER)
        _set_cell(ws, row, 2, "내용", fill=FILL_HEADER, font=FONT_HEADER, align=ALIGN_CENTER)
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=8)
        _set_cell(ws, row, 3, "근거", fill=FILL_HEADER, font=FONT_HEADER, align=ALIGN_CENTER)
        row += 1

        items = section.get("items", []) or []
        used_names = set()

        for item in items[:8]:
            name = clean_text(item.get("name", ""))
            value = clean_text(item.get("value", ""))
            evidence = [clean_text(x) for x in (item.get("evidence", []) or []) if clean_text(x)]

            if not name or name in used_names:
                continue
            used_names.add(name)

            _set_cell(ws, row, 1, name)
            _set_cell(ws, row, 2, value or "공개 자료 기준 직접 확인 제한")
            ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=8)
            _set_cell(ws, row, 3, " / ".join(evidence[:3]) if evidence else "공개 자료 기준 직접 확인 제한")
            row += 1

        # 템플릿 필드 보정
        for field_name in fields:
            if field_name not in used_names:
                _set_cell(ws, row, 1, field_name)
                _set_cell(ws, row, 2, "공개 자료 기준 직접 확인 제한")
                ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=8)
                _set_cell(ws, row, 3, "관련 근거 문구 직접 확인 제한")
                row += 1

        row += 1


def _write_quant_sheet(ws, company: dict, categories: list[dict], score_payload: dict):
    corp_name = clean_text(company.get("corp_name", "기업"))
    score_lookup = _score_map(score_payload)

    ws.merge_cells("A1:H1")
    _set_cell(
        ws,
        1,
        1,
        f"{corp_name} 정량지표 정리",
        fill=FILL_HEADER,
        font=FONT_TITLE,
        align=ALIGN_CENTER,
    )

    headers = [
        "섹션",
        "평가축",
        "점수",
        "정량 포인트",
        "정량 포인트 수",
        "원문 수치 신호",
        "근거 신호 수",
        "비고",
    ]

    row = 3
    for idx, header in enumerate(headers, start=1):
        _set_cell(ws, row, idx, header, fill=FILL_HEADER, font=FONT_HEADER, align=ALIGN_CENTER)

    row += 1
    for section_name in CANONICAL_SECTION_ORDER:
        section = _find_section(categories, section_name)
        linked_score = score_lookup.get(section_name, {})

        quant_points = [clean_text(x) for x in (section.get("quant_points", []) or []) if clean_text(x)]
        result_text = clean_text(section.get("result", ""))

        numeric_signals = []
        for token in _extract_inline_numeric_signals(result_text):
            numeric_signals.append(token)

        for item in section.get("items", []) or []:
            for token in _extract_inline_numeric_signals(clean_text(item.get("value", ""))):
                numeric_signals.append(token)
            for ev in item.get("evidence", []) or []:
                for token in _extract_inline_numeric_signals(clean_text(ev)):
                    numeric_signals.append(token)

        numeric_signals = _uniq(numeric_signals)

        _set_cell(ws, row, 1, section_name)
        _set_cell(ws, row, 2, clean_text(linked_score.get("axis", "-")))
        _set_cell(ws, row, 3, linked_score.get("score", 0), align=ALIGN_CENTER)
        _set_cell(ws, row, 4, " / ".join(quant_points) if quant_points else "정량 포인트 직접 확인 제한")
        _set_cell(ws, row, 5, len(quant_points), align=ALIGN_CENTER)
        _set_cell(ws, row, 6, " / ".join(numeric_signals[:10]) if numeric_signals else "-")
        _set_cell(ws, row, 7, linked_score.get("evidence_count", 0), align=ALIGN_CENTER)

        note = "근거 부족 시 보수 반영"
        if linked_score.get("score", 0) >= 4:
            note = "정량/근거 신호 비교적 충분"
        elif linked_score.get("score", 0) <= 2:
            note = "근거 부족 또는 정량 신호 제한"
        _set_cell(ws, row, 8, note)

        row += 1


def _write_source_sheet(ws, company: dict):
    corp_name = clean_text(company.get("corp_name", "기업"))
    ws.merge_cells("A1:H1")
    _set_cell(
        ws,
        1,
        1,
        f"{corp_name} 원천 문서/URL 정리",
        fill=FILL_HEADER,
        font=FONT_TITLE,
        align=ALIGN_CENTER,
    )

    headers = ["구분", "버킷", "URL/내용", "비고"]
    row = 3
    for idx, h in enumerate(headers, start=1):
        _set_cell(ws, row, idx, h, fill=FILL_HEADER, font=FONT_HEADER, align=ALIGN_CENTER)

    row += 1

    # homepage / ir
    single_fields = [
        ("단일URL", "homepage_url", clean_text(company.get("homepage_url", ""))),
        ("단일URL", "ir_url", clean_text(company.get("ir_url", ""))),
    ]
    for kind, bucket, value in single_fields:
        if value:
            _set_cell(ws, row, 1, kind)
            _set_cell(ws, row, 2, bucket)
            _set_cell(ws, row, 3, value)
            _set_cell(ws, row, 4, "-")
            row += 1

    # urls
    url_map = company.get("urls", {}) or {}
    if isinstance(url_map, dict):
        for bucket, value in url_map.items():
            if isinstance(value, list):
                for v in value:
                    vv = clean_text(v)
                    if not vv:
                        continue
                    _set_cell(ws, row, 1, "urls")
                    _set_cell(ws, row, 2, clean_text(bucket))
                    _set_cell(ws, row, 3, vv)
                    _set_cell(ws, row, 4, "-")
                    row += 1
            else:
                vv = clean_text(value)
                if vv:
                    _set_cell(ws, row, 1, "urls")
                    _set_cell(ws, row, 2, clean_text(bucket))
                    _set_cell(ws, row, 3, vv)
                    _set_cell(ws, row, 4, "-")
                    row += 1

    # extra_urls 구버전 호환
    extra_url_map = company.get("extra_urls", {}) or {}
    if isinstance(extra_url_map, dict):
        for bucket, value in extra_url_map.items():
            if isinstance(value, list):
                for v in value:
                    vv = clean_text(v)
                    if not vv:
                        continue
                    _set_cell(ws, row, 1, "extra_urls")
                    _set_cell(ws, row, 2, clean_text(bucket))
                    _set_cell(ws, row, 3, vv)
                    _set_cell(ws, row, 4, "구버전 호환")
                    row += 1
            else:
                vv = clean_text(value)
                if vv:
                    _set_cell(ws, row, 1, "extra_urls")
                    _set_cell(ws, row, 2, clean_text(bucket))
                    _set_cell(ws, row, 3, vv)
                    _set_cell(ws, row, 4, "구버전 호환")
                    row += 1

    # notes / text
    note_fields = ["notes", "report_text", "homepage_text", "ir_text", "research_text", "seed_text"]
    for field in note_fields:
        value = clean_text(company.get(field, ""))
        if value:
            _set_cell(ws, row, 1, "text")
            _set_cell(ws, row, 2, field)
            _set_cell(ws, row, 3, value)
            _set_cell(ws, row, 4, "원문 텍스트")
            row += 1


def _extract_inline_numeric_signals(text: str) -> list[str]:
    text = clean_text(text)
    if not text:
        return []

    tokens = []
    for m in __import__("re").findall(r"\d+(?:[\.,]\d+)?\s*(?:%|배|건|개|회|종|명|억원|천원|백만원)?", text):
        token = clean_text(m)
        if token:
            tokens.append(token)

    return _uniq(tokens)


def _uniq(values: list[str]) -> list[str]:
    uniq = []
    seen = set()
    for v in values:
        vv = clean_text(v)
        if not vv or vv in seen:
            continue
        seen.add(vv)
        uniq.append(vv)
    return uniq