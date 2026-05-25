from __future__ import annotations

from functools import lru_cache
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import load_workbook

from .config import DEFAULT_TEMPLATE_PATH, TECH_ROW_SCHEMA_PATH
from .utils import clean_text, load_json


def _safe_load_wb(path: str | Path):
    path = Path(path)
    if not path.exists():
        return None
    return load_workbook(path)


@lru_cache(maxsize=4)
def load_template_workbook(template_path: str | Path = DEFAULT_TEMPLATE_PATH):
    return _safe_load_wb(template_path)


@lru_cache(maxsize=4)
def load_row_schema_json() -> list[dict]:
    rows = load_json(TECH_ROW_SCHEMA_PATH, default=[])
    return rows or []


@lru_cache(maxsize=8)
def load_formula_map(template_path: str | Path = DEFAULT_TEMPLATE_PATH) -> Dict[str, List[dict]]:
    wb = load_template_workbook(template_path)
    out: Dict[str, List[dict]] = {}
    if wb is None or "수식 정리" not in wb.sheetnames:
        return out

    ws = wb["수식 정리"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        values = [clean_text(v or "") for v in row[:6]]
        if not any(values):
            continue
        category = values[0]
        metric_name = values[1]
        formula = values[2]
        description = values[3] if len(values) > 3 else ""
        source_hint = values[4] if len(values) > 4 else ""
        example = values[5] if len(values) > 5 else ""
        if not category or not metric_name:
            continue
        out.setdefault(category, []).append(
            {
                "metric_name": metric_name,
                "formula": formula,
                "description": description,
                "source_hint": source_hint,
                "example": example,
            }
        )
    return out


@lru_cache(maxsize=8)
def load_base_sheet_map(template_path: str | Path = DEFAULT_TEMPLATE_PATH) -> Dict[str, dict]:
    wb = load_template_workbook(template_path)
    out: Dict[str, dict] = {}
    if wb is None or "베이스" not in wb.sheetnames:
        return out

    ws = wb["베이스"]
    current_category = ""
    for r in range(2, ws.max_row + 1):
        category = clean_text(ws.cell(r, 2).value or "")
        major_source = clean_text(ws.cell(r, 3).value or "")
        sub_source = clean_text(ws.cell(r, 4).value or "")
        quant_parts = clean_text(ws.cell(r, 5).value or "")
        item_name = clean_text(ws.cell(r, 6).value or "")
        result_seed = clean_text(ws.cell(r, 7).value or "")

        if category:
            current_category = category
            out.setdefault(
                current_category,
                {
                    "category": current_category,
                    "items": [],
                    "major_source": major_source,
                    "sub_source": sub_source,
                    "quant_parts": quant_parts,
                    "result_template": result_seed,
                },
            )

        if not current_category:
            continue

        if item_name and item_name not in out[current_category]["items"]:
            out[current_category]["items"].append(item_name)

        if result_seed and not out[current_category].get("result_template"):
            out[current_category]["result_template"] = result_seed

    return out


@lru_cache(maxsize=8)
def load_company_sheet_examples(company_name: str, template_path: str | Path = DEFAULT_TEMPLATE_PATH) -> Dict[str, dict]:
    wb = load_template_workbook(template_path)
    if wb is None:
        return {}

    norm_company = _norm(company_name)
    target_ws = None

    for ws_name in wb.sheetnames:
        if ws_name in {"베이스", "수식 정리", "요약", "상세분석", "점수", "정량지표"}:
            continue
        ws = wb[ws_name]
        a2 = clean_text(ws["A2"].value or "")
        candidates = {_norm(ws_name), _norm(a2)}
        if norm_company in candidates or any(norm_company and norm_company in c for c in candidates):
            target_ws = ws
            break

    if target_ws is None:
        return {}

    current_category = ""
    out: Dict[str, dict] = {}
    for r in range(2, target_ws.max_row + 1):
        category = clean_text(target_ws.cell(r, 2).value or "")
        item_name = clean_text(target_ws.cell(r, 6).value or "")
        content = clean_text(target_ws.cell(r, 7).value or "")

        if category:
            current_category = category
            out.setdefault(current_category, {"summary": "", "items": {}})

        if not current_category:
            continue

        if item_name and content:
            out[current_category]["items"][item_name] = content
        elif content and not out[current_category]["summary"]:
            out[current_category]["summary"] = content

    return out


@lru_cache(maxsize=8)
def get_template_rows(template_path: str | Path = DEFAULT_TEMPLATE_PATH) -> List[dict]:
    json_rows = load_row_schema_json()
    base_map = load_base_sheet_map(template_path)
    formula_map = load_formula_map(template_path)

    rows: List[dict] = []
    seen = set()

    for row in json_rows:
        category = clean_text(row.get("category", ""))
        if not category:
            continue
        merged = dict(row)
        if category in base_map:
            merged.setdefault("major_source", base_map[category].get("major_source", ""))
            merged.setdefault("sub_source", base_map[category].get("sub_source", ""))
            merged.setdefault("result_template", row.get("result_template", "") or base_map[category].get("result_template", ""))
            if not merged.get("items"):
                merged["items"] = base_map[category].get("items", [])
        merged["formula_metrics"] = formula_map.get(category, [])
        rows.append(merged)
        seen.add(category)

    for category, info in base_map.items():
        if category in seen:
            continue
        rows.append(
            {
                "category": category,
                "items": info.get("items", []),
                "major_source": info.get("major_source", ""),
                "sub_source": info.get("sub_source", ""),
                "quant_rule": info.get("quant_parts", ""),
                "result_template": info.get("result_template", ""),
                "formula_metrics": formula_map.get(category, []),
            }
        )

    return rows


@lru_cache(maxsize=2)
def get_category_row_lookup(template_path: str | Path = DEFAULT_TEMPLATE_PATH) -> Dict[str, dict]:
    wb = load_template_workbook(template_path)
    if wb is None or "베이스" not in wb.sheetnames:
        return {}
    ws = wb["베이스"]
    current_category = ""
    out: Dict[str, dict] = {}

    for r in range(2, ws.max_row + 1):
        category = clean_text(ws.cell(r, 2).value or "")
        item_name = clean_text(ws.cell(r, 6).value or "")

        if category:
            current_category = category
            out.setdefault(current_category, {"summary_row": r, "item_rows": {}})

        if current_category and item_name:
            out[current_category]["item_rows"][item_name] = r

    return out


def _norm(text: str) -> str:
    return (
        clean_text(text)
        .replace("주식회사", "")
        .replace("(주)", "")
        .replace("㈜", "")
        .replace(" ", "")
        .lower()
    )