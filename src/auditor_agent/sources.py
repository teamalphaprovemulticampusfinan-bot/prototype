from __future__ import annotations

import csv
import html as html_lib
import json
import os
import re
from pathlib import Path
from common.data_paths import DATA_DIR, company_common_dir, company_root, field_common_dir
from typing import Any
from urllib.parse import quote_plus

import requests
import yaml

ROOT = Path(__file__).resolve().parents[2]
WORKSPACE = DATA_DIR


def _env_bool(name: str, default: bool = False) -> bool:
    value = os.getenv(name)
    if value is None or str(value).strip() == "":
        return default
    return str(value).strip().lower() in {"1", "true", "yes", "y", "on"}


def _strip_html(text: Any) -> str:
    if text is None:
        return ""
    cleaned = re.sub(r"<[^>]+>", " ", str(text))
    cleaned = html_lib.unescape(cleaned)
    return re.sub(r"\s+", " ", cleaned).strip()


def _read_text(path: Path, limit: int = 6000) -> str:
    try:
        txt = path.read_text(encoding="utf-8", errors="ignore")
        return txt[:limit]
    except Exception:
        return ""


def _yaml_context(path: Path) -> str:
    try:
        data = yaml.safe_load(path.read_text(encoding="utf-8", errors="ignore")) or {}
        return json.dumps(data, ensure_ascii=False, indent=2)[:8000]
    except Exception:
        return _read_text(path, 4000)


def _csv_context(path: Path, rows: int = 35, chars: int = 7000) -> str:
    try:
        out: list[str] = []
        with path.open("r", encoding="utf-8-sig", errors="ignore", newline="") as fp:
            reader = csv.reader(fp)
            for idx, row in enumerate(reader):
                if idx >= rows:
                    break
                out.append(" | ".join(str(x)[:120] for x in row[:18]))
        return "\n".join(out)[:chars]
    except Exception:
        return _read_text(path, chars)


def _xlsx_context(path: Path, company: str, rows: int = 40, chars: int = 7000) -> str:
    try:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        chunks: list[str] = []
        company_re = re.compile(re.escape(company), re.I) if company else None

        for ws in wb.worksheets[:5]:
            lines: list[str] = []
            for ridx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                if ridx > 250:
                    break

                vals = ["" if v is None else str(v) for v in row[:12]]
                text = " | ".join(v[:120] for v in vals if v.strip())
                if not text:
                    continue

                if ridx <= 5 or not company_re or company_re.search(text):
                    lines.append(text)

                if len(lines) >= rows:
                    break

            if lines:
                chunks.append(f"[{path.name}:{ws.title}]\n" + "\n".join(lines))

        return "\n\n".join(chunks)[:chars]
    except Exception:
        return ""


def _naver_news_snippets(
    query: str,
    timeout: int = 8,
    max_results: int = 6,
) -> list[dict[str, str]]:
    """
    Auditor 웹 교차검증용 네이버 뉴스 검색입니다.

    필요 환경변수:
    - AUDITOR_FIRST_ENABLE_WEB=1
    - AUDITOR_WEB_ENABLE_NAVER=1  # 생략 시 기본 1
    - NAVER_CLIENT_ID=...
    - NAVER_CLIENT_SECRET=...

    네이버 호출에 실패하면 빈 리스트를 반환하고, 이후 Serper/DuckDuckGo fallback으로 넘어갑니다.
    """
    if not _env_bool("AUDITOR_WEB_ENABLE_NAVER", True):
        return []

    client_id = os.getenv("NAVER_CLIENT_ID", "").strip()
    client_secret = os.getenv("NAVER_CLIENT_SECRET", "").strip()
    if not client_id or not client_secret:
        return []

    try:
        response = requests.get(
            "https://openapi.naver.com/v1/search/news.json",
            headers={
                "X-Naver-Client-Id": client_id,
                "X-Naver-Client-Secret": client_secret,
            },
            params={
                "query": query,
                "display": max(1, min(int(max_results), 100)),
                "start": 1,
                "sort": "date",
            },
            timeout=timeout,
        )
        response.raise_for_status()
        data = response.json()
    except Exception:
        return []

    items: list[dict[str, str]] = []
    for block in data.get("items", [])[:max_results]:
        if not isinstance(block, dict):
            continue
        title = _strip_html(block.get("title"))
        snippet = _strip_html(block.get("description"))
        source = str(block.get("originallink") or block.get("link") or "naver_news").strip()
        if title or snippet:
            items.append({"title": title, "snippet": snippet, "source": source})

    return items[:max_results]


def _serper_snippets(
    query: str,
    timeout: int = 8,
    max_results: int = 6,
) -> list[dict[str, str]]:
    """
    Serper fallback 검색입니다.

    현재 기본 경로는 네이버 뉴스입니다. Serper까지 쓰려면 SERPER_API_KEY 또는 GOOGLE_SERPER_API_KEY를 넣으면 됩니다.
    API 제한/429가 걱정되면 SERPER_API_KEY를 비워두거나 AUDITOR_FIRST_ENABLE_WEB=0으로 두면 됩니다.
    """
    api_key = os.getenv("SERPER_API_KEY") or os.getenv("GOOGLE_SERPER_API_KEY")
    if not api_key:
        return []

    url = "https://google.serper.dev/search"
    headers = {
        "X-API-KEY": api_key,
        "Content-Type": "application/json",
    }
    payload = {
        "q": query,
        "gl": "kr",
        "hl": "ko",
        "num": max_results,
    }

    try:
        response = requests.post(url, headers=headers, json=payload, timeout=timeout)
        response.raise_for_status()
        data = response.json()

        items: list[dict[str, str]] = []

        for block in data.get("organic", [])[:max_results]:
            if not isinstance(block, dict):
                continue

            title = str(block.get("title") or "").strip()
            snippet = str(block.get("snippet") or "").strip()
            source = str(block.get("link") or "serper").strip()

            if title or snippet:
                items.append({"title": title, "snippet": snippet, "source": source})

        for block in data.get("news", [])[:max_results]:
            if not isinstance(block, dict):
                continue

            title = str(block.get("title") or "").strip()
            snippet = str(block.get("snippet") or "").strip()
            source = str(block.get("link") or block.get("source") or "serper_news").strip()

            if title or snippet:
                items.append({"title": title, "snippet": snippet, "source": source})

        return items[:max_results]

    except Exception:
        return []


# NewsAPI는 감사 웹 검증에는 기본 사용하지 않습니다.
# 나중에 NewsAPI로도 Auditor 웹 검증을 추가하고 싶으면 아래 흐름으로 별도 함수를 만들면 됩니다.
# def _newsapi_snippets(query: str, timeout: int = 8, max_results: int = 6) -> list[dict[str, str]]:
#     if not _env_bool("AUDITOR_WEB_ENABLE_NEWS_API", False):
#         return []
#     api_key = os.getenv("NEWS_API_KEY", "").strip()
#     if not api_key:
#         return []
#     ...


def _duckduckgo_snippets(
    query: str,
    timeout: int = 8,
    max_results: int = 6,
) -> list[dict[str, str]]:
    url = "https://duckduckgo.com/html/?q=" + quote_plus(query)
    headers = {"User-Agent": "Mozilla/5.0 (compatible; FirstAuditor/1.0)"}

    try:
        response = requests.get(url, headers=headers, timeout=timeout)
        response.raise_for_status()
        page_html = response.text

        items: list[dict[str, str]] = []

        blocks = re.split(r'<div class="result__body">', page_html)[1:]
        for block in blocks[:max_results]:
            title_match = re.search(r'<a[^>]+class="result__a"[^>]*>(.*?)</a>', block, re.S)
            snippet_match = re.search(r'<a[^>]+class="result__snippet"[^>]*>(.*?)</a>', block, re.S)
            href_match = re.search(r'<a[^>]+class="result__a"[^>]+href="([^"]+)"', block, re.S)

            title = re.sub(r"<.*?>", "", title_match.group(1)) if title_match else ""
            snippet = re.sub(r"<.*?>", "", snippet_match.group(1)) if snippet_match else ""
            source = href_match.group(1) if href_match else "duckduckgo"

            clean = re.sub(r"\s+", " ", f"{title} {snippet}").strip()
            if clean:
                items.append({"title": title.strip(), "snippet": snippet.strip(), "source": source})

        return items

    except Exception:
        return []


def _public_search_snippets(
    query: str,
    timeout: int = 8,
    max_results: int = 6,
) -> list[dict[str, str]]:
    """네이버 뉴스 우선, Serper fallback, DuckDuckGo 최종 fallback."""
    items = _naver_news_snippets(query=query, timeout=timeout, max_results=max_results)
    if items:
        return items[:max_results]

    items = _serper_snippets(query=query, timeout=timeout, max_results=max_results)
    if items:
        return items[:max_results]

    return _duckduckgo_snippets(query=query, timeout=timeout, max_results=max_results)[:max_results]


def _dedupe_web_items(items: list[dict[str, str]], limit: int = 12) -> list[dict[str, str]]:
    seen: set[str] = set()
    deduped: list[dict[str, str]] = []

    for item in items:
        title = str(item.get("title") or "").strip()
        snippet = str(item.get("snippet") or "").strip()
        source = str(item.get("source") or "").strip()

        key = re.sub(r"\s+", " ", f"{title} {snippet}")[:250]
        if not key or key in seen:
            continue

        seen.add(key)
        deduped.append({"title": title, "snippet": snippet, "source": source})

        if len(deduped) >= limit:
            break

    return deduped


def collect_source_context(
    company_dir: str,
    company: str,
    max_chars: int = 24000,
    enable_web: bool = True,
) -> dict[str, Any]:
    """Collect non-secret evidence for first-stage factual checking.

    이 함수는 .env 파일을 읽지 않습니다.
    환경변수는 이미 런타임에 로드된 값을 os.getenv로만 참조합니다.
    """
    company_path = company_common_dir(company_dir)
    company_base = company_root(company_dir)
    chunks: list[str] = []
    source_files: list[str] = []

    yaml_path = company_path / "company.yaml"
    if yaml_path.exists():
        chunks.append(f"[company.yaml]\n{_yaml_context(yaml_path)}")
        source_files.append(str(yaml_path))

    scan_dirs = [
        company_path,
        company_base / "finance",
        company_base / "market",
        company_base / "tech",
        company_base / "issue",
        company_base / "macro",
    ]

    for pattern in ("*.csv", "*.txt", "*.json", "*.md"):
        for scan_dir in scan_dirs:
            if not scan_dir.exists():
                continue
            for path in sorted(scan_dir.glob(pattern))[:8]:
                if path.name == ".env" or path.suffix.lower() == ".env":
                    continue

                if path.suffix.lower() == ".csv":
                    text = _csv_context(path)
                else:
                    text = _read_text(path, 6000)

                if text.strip():
                    chunks.append(f"[{path.name}]\n{text}")
                    source_files.append(str(path))

    for scan_dir in scan_dirs:
        if not scan_dir.exists():
            continue
        for path in sorted(scan_dir.glob("*.xlsx"))[:5]:
            text = _xlsx_context(path, company)
            if text.strip():
                chunks.append(text)
                source_files.append(str(path))

    for path in sorted((field_common_dir("data")).glob("*.xlsx"))[:8]:
        text = _xlsx_context(path, company, rows=25, chars=5000)
        if text.strip():
            chunks.append(text)
            source_files.append(str(path))

    web_items: list[dict[str, str]] = []

    if enable_web:
        timeout = int(os.getenv("REQUEST_TIMEOUT", "25") or 25)

        queries = [
            f"{company} 기업 개요",
            f"{company} 최근 뉴스",
            f"{company} 실적 사업 반도체",
            f"{company} 수주 계약 고객사",
            f"{company} 투자위험 리스크",
        ]

        for query in queries:
            web_items.extend(
                _public_search_snippets(
                    query=query,
                    timeout=min(timeout, 10),
                    max_results=4,
                )
            )

        web_items = _dedupe_web_items(web_items, limit=14)

        if web_items:
            web_text = "\n".join(
                f"- {item.get('title', '')} :: {item.get('snippet', '')} :: {item.get('source', '')}"
                for item in web_items
            )
            chunks.append(f"[public_search_snippets]\n{web_text}")

    context = "\n\n".join(chunks)

    return {
        "context": context[:max_chars],
        "source_files": source_files,
        "web_items": web_items,
        "has_context": bool(context.strip()),
    }
